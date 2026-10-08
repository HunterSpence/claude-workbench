import importlib.util, os, shutil, sys, tempfile, time, unittest, zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import doc_index as di

HAVE = {m: importlib.util.find_spec(m) is not None for m in ("markitdown", "docx", "pptx", "openpyxl")}
LABEL_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
             'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
             '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" '
             'name="MSIP_Label_00000000-0000-0000-0000-000000000000_Name"><vt:lpwstr>Internal</vt:lpwstr></property></Properties>')


GUID = "11111111-2222-3333-4444-555555555555"
TENANT = "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee"


def label_info(enabled="1", removed="0"):
    """A LabelInfo part in the [MS-OFFCRYPTO] 2.6.4 shape (namespace, labelList, label attributes)."""
    return ('<?xml version="1.0" encoding="utf-8" standalone="yes"?>'
            '<clbl:labelList xmlns:clbl="http://schemas.microsoft.com/office/2020/mipLabelMetadata">'
            f'<clbl:label id="{{{GUID}}}" enabled="{enabled}" method="Standard" siteId="{{{TENANT}}}" '
            f'contentBits="0" removed="{removed}"/></clbl:labelList>')


def quiet(*_):
    pass


class DocIndexTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.root = os.path.join(self.tmp, "library")
        os.makedirs(os.path.join(self.root, "sub"))
        self.db = os.path.join(self.tmp, "idx", "index.sqlite")
        self.w("notes.md", "# Onboarding\n\nThe quarterly rollout plan covers three regions.\n\nSecond paragraph about budgets.\n")
        self.w("sub/readme.txt", "Plain text mentioning the escalation matrix.\n")
        with open(os.path.join(self.root, "locked.docx"), "wb") as f:
            f.write(di.OLE_MAGIC + b"\0" * 100)
        with open(os.path.join(self.root, "~$notes.docx"), "wb") as f:
            f.write(b"owner lock file")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def w(self, rel, text):
        with open(os.path.join(self.root, rel), "w", encoding="utf-8") as f:
            f.write(text)

    def run_index(self, **kw):
        return di.index(self.root, self.db, log=quiet, **kw)

    def test_index_search_incremental_rebuild(self):
        s = self.run_index()
        self.assertEqual((s["new"], s["failed"]), (2, 1))
        hits = di.search(self.db, "rollout regions")
        self.assertEqual(hits[0]["ref"], "notes.md#p2")
        self.assertEqual(hits[0]["heading"], "Onboarding")
        self.assertIn("[rollout]", hits[0]["snippet"])
        self.assertEqual(di.search(self.db, "escalation")[0]["ref"], "sub/readme.txt#p1")
        self.assertEqual(di.search(self.db, "rollout zebra")[0]["ref"], "notes.md#p2")   # OR fallback
        self.assertIn(">> notes.md#p2", di.show(self.db, "notes.md#p2", 1))
        failed = {p: (s, e) for p, s, e in di.report(self.db)}
        self.assertEqual(failed["locked.docx"][0], "failed")
        self.assertIn("encrypted", failed["locked.docx"][1])
        self.assertNotIn("~$notes.docx", failed)

        s = self.run_index()                       # nothing changed
        self.assertEqual((s["new"], s["updated"], s["unchanged"]), (0, 0, 3))

        time.sleep(0.05)
        self.w("notes.md", "# Onboarding\n\nThe plan was replaced by a pilot.\n")
        s = self.run_index()
        self.assertEqual(s["updated"], 1)
        self.assertEqual(di.search(self.db, "rollout"), [])
        self.assertEqual(di.search(self.db, "pilot")[0]["ref"], "notes.md#p2")

        os.remove(os.path.join(self.root, "sub", "readme.txt"))
        s = self.run_index()                       # incremental never deletes
        self.assertEqual(s["removed"], 0)
        self.assertTrue(di.search(self.db, "escalation"))
        s = self.run_index(rebuild=True)
        self.assertEqual(s["removed"], 1)
        self.assertEqual(di.search(self.db, "escalation"), [])

    def test_rebuild_refuses_to_prune_an_empty_folder(self):
        self.run_index()
        for name in os.listdir(self.root):
            p = os.path.join(self.root, name)
            shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
        s = self.run_index(rebuild=True)
        self.assertEqual(s["removed"], 0)
        self.assertTrue(di.search(self.db, "rollout"))

    def test_label_reader_both_locations(self):
        def z(name, parts):
            p = os.path.join(self.tmp, name)
            with zipfile.ZipFile(p, "w") as f:
                for k, v in parts.items():
                    f.writestr(k, v)
            return p
        self.assertEqual(di.sensitivity_label(z("a.docx", {"docMetadata/LabelInfo.xml": label_info()})), f"label {{{GUID}}}")
        self.assertIsNone(di.sensitivity_label(z("b.docx", {"docMetadata/LabelInfo.xml": label_info(removed="1")})))
        self.assertIsNone(di.sensitivity_label(z("c.docx", {"docMetadata/LabelInfo.xml": label_info(enabled="0")})))
        self.assertEqual(di.sensitivity_label(z("d.docx", {"docProps/custom.xml": LABEL_XML})), "Internal")
        self.assertEqual(di.sensitivity_label(z("e.docx", {"docProps/custom.xml": LABEL_XML,
                                                          "docMetadata/LabelInfo.xml": label_info()})), "Internal")
        self.assertIsNone(di.sensitivity_label(z("f.docx", {"word/document.xml": "<w/>"})))          # no label parts: fine
        self.assertTrue(di.sensitivity_label(z("g.docx", {"docMetadata/LabelInfo.xml": "<not xml"})))  # fails closed
        self.assertTrue(di.sensitivity_label(z("h.docx", {"DOCMETADATA/labelinfo.XML": label_info()})))  # case-insensitive
        self.assertIsNone(di.sensitivity_label(z("i.docx", {"other/LabelInfo.xml": label_info()})))      # exact path only

    def test_oversized_label_parts_fail_closed_without_being_read(self):
        big = " " * (di.LABEL_PART_CAP + 10)          # compresses to almost nothing: a tiny zip bomb
        opened = []
        real_open = zipfile.ZipFile.open
        def spy(zf, member, *a, **k):
            opened.append(member.filename if isinstance(member, zipfile.ZipInfo) else member)
            return real_open(zf, member, *a, **k)
        for part in ("docMetadata/LabelInfo.xml", "docProps/custom.xml"):
            p = os.path.join(self.tmp, "big.docx")
            with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as f:
                f.writestr(part, big)
            self.assertLess(os.path.getsize(p), 20000)
            zipfile.ZipFile.open = spy
            try:
                label = di.sensitivity_label(p)
            finally:
                zipfile.ZipFile.open = real_open
            self.assertTrue(label and "oversized" in label, part)
            self.assertNotIn(part, opened, f"{part} was read despite its size")

    def test_zip_bomb_office_file_is_skipped_before_conversion(self):
        p = os.path.join(self.root, "bomb.docx")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as f:
            f.writestr("word/document.xml", " " * (3 * 1024 * 1024))   # 3 MB that deflates to a few KB
        converted = []
        orig = di.to_markdown
        di.to_markdown = lambda path: converted.append(path) or orig(path)
        try:
            self.run_index(max_unzipped_mb=2)
        finally:
            di.to_markdown = orig
        self.assertFalse([c for c in converted if c.endswith("bomb.docx")])
        rows = {p_: (st, e) for p_, st, e in di.report(self.db)}
        self.assertEqual(rows["bomb.docx"][0], "skipped")
        self.assertIn("too large when uncompressed", rows["bomb.docx"][1])

    def test_purge_deletes_index_and_plaintext_cache(self):
        self.run_index()
        cache = os.path.join(os.path.dirname(self.db), "cache")
        self.assertTrue(os.listdir(cache))
        self.assertEqual(di.main(["--db", self.db, "--purge"]), 0)
        self.assertFalse(os.path.exists(self.db))
        self.assertFalse(os.path.exists(cache))
        self.assertTrue(os.path.exists(os.path.join(self.root, "notes.md")))   # originals untouched

    def test_cloud_only_placeholder_is_never_opened(self):
        with open(os.path.join(self.root, "online-only.md"), "w") as f:
            f.write("x" * 7777)
        opened = []
        orig_sha, orig_md, orig_cloud = di.sha256, di.to_markdown, di.is_cloud_only
        di.sha256 = lambda p: opened.append(p) or orig_sha(p)
        di.to_markdown = lambda p: opened.append(p) or orig_md(p)
        di.is_cloud_only = lambda st: st.st_size == 7777
        try:
            s = self.run_index()
        finally:
            di.sha256, di.to_markdown, di.is_cloud_only = orig_sha, orig_md, orig_cloud
        self.assertEqual(s["cloud_only"], 1)
        self.assertFalse([p for p in opened if p.endswith("online-only.md")])
        self.assertIn(("online-only.md", "cloud-only"), [(p, st) for p, st, _ in di.report(self.db)])

    @unittest.skipUnless(all(HAVE.values()), "needs markitdown, python-docx, python-pptx, openpyxl")
    def test_office_formats_and_label(self):
        import docx, openpyxl, pptx
        d = docx.Document()
        d.add_heading("Vendor review", 1)
        d.add_paragraph("The procurement checklist requires two signatures.")
        dp = os.path.join(self.root, "review.docx")
        d.save(dp)
        with zipfile.ZipFile(dp, "a") as z:
            z.writestr("docProps/custom.xml", LABEL_XML)
        prs = pptx.Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Roadmap"
        s.placeholders[1].text = "Milestone telemetry dashboard ships in spring"
        prs.save(os.path.join(self.root, "deck.pptx"))
        wb = openpyxl.Workbook()
        wb.active.append(["metric", "value"])
        wb.active.append(["throughput", 42])
        wb.save(os.path.join(self.root, "numbers.xlsx"))
        self.w("page.html", "<html><body><h1>Policy</h1><p>Retention period is seven years.</p></body></html>")
        st = self.run_index()
        self.assertEqual((st["failed"], st["labelled"]), (1, 1), di.report(self.db))   # the OLE file; the labelled docx
        self.assertEqual(di.search(self.db, "procurement signatures"), [])               # labelled content not indexed
        self.assertIn(("review.docx", "labelled"), [(p, s_) for p, s_, _ in di.report(self.db)])
        self.assertEqual(self.run_index()["labelled"], 1)                                 # stays skipped, cheaply
        self.run_index(include_labelled=True)                                             # explicit opt-in
        hit = di.search(self.db, "procurement signatures")[0]
        self.assertTrue(hit["ref"].startswith("review.docx#p"))
        self.assertEqual(hit["label"], "Internal")
        self.assertTrue(di.search(self.db, "telemetry dashboard")[0]["ref"].startswith("deck.pptx#p"))
        self.assertTrue(di.search(self.db, "throughput")[0]["ref"].startswith("numbers.xlsx#p"))
        self.assertTrue(di.search(self.db, "retention")[0]["ref"].startswith("page.html#p"))

    @unittest.skipUnless(HAVE["markitdown"] and HAVE["docx"], "needs markitdown and python-docx")
    def test_labelinfo_only_docx_is_skipped_unless_opted_in(self):
        import docx
        d = docx.Document()
        d.add_paragraph("Coauthoring budget forecast for the northern programme.")
        p = os.path.join(self.root, "coauthored.docx")
        d.save(p)
        with zipfile.ZipFile(p, "a") as z:                         # label ONLY in the co-authoring location
            z.writestr("docMetadata/LabelInfo.xml", label_info())
        with zipfile.ZipFile(p) as z:
            self.assertNotIn("docProps/custom.xml", z.namelist())
        st = self.run_index()
        self.assertEqual(st["labelled"], 1, di.report(self.db))
        self.assertEqual(di.search(self.db, "coauthoring forecast"), [])
        self.run_index(include_labelled=True)
        hit = di.search(self.db, "coauthoring forecast")[0]
        self.assertTrue(hit["ref"].startswith("coauthored.docx#p"))
        self.assertEqual(hit["label"], f"label {{{GUID}}}")


if __name__ == "__main__":
    unittest.main()

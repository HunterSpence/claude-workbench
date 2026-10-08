import importlib.util, json, os, shutil, sys, tempfile, unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


@unittest.skipUnless(importlib.util.find_spec("pptx"), "needs python-pptx")
class BuildDeckTest(unittest.TestCase):
    def setUp(self):
        import build_deck
        from PIL import Image
        self.bd = build_deck
        self.d = tempfile.mkdtemp()
        Image.new("RGB", (400, 200), (30, 90, 160)).save(os.path.join(self.d, "pic.png"))
        self.outline = {"slides": [
            {"title": "Quarterly review", "subtitle": "Platform team", "notes": "Open with the headline."},
            {"title": "What changed", "bullets": ["Faster builds", ["cache hits up", "fewer retries"], "New owner"]},
            {"title": "Costs", "table": {"columns": ["Item", "Q2", "Q3"], "rows": [["Compute", 12, 9], ["Storage", 3, None]]}},
            {"title": "Trend", "chart": {"type": "line", "categories": ["Jan", "Feb", "Mar"], "series": {"Users": [10, 14, 19]},
                                         "data_labels": True, "number_format": "0"}},
            {"title": "Diagram", "image": "pic.png"},
            {"title": "x" * 90, "bullets": [str(i) for i in range(9)]},
        ]}

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def build(self, outline):
        p = os.path.join(self.d, "outline.json")
        with open(p, "w") as f:
            json.dump(outline, f)
        out = os.path.join(self.d, "deck.pptx")
        self.assertEqual(self.bd.main(["build", p, out]), 0)
        from pptx import Presentation
        return Presentation(out)

    def test_every_slide_kind(self):
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = self.build(self.outline)
        s = list(prs.slides)
        self.assertEqual(len(s), 6)
        self.assertEqual(s[0].shapes.title.text, "Quarterly review")
        self.assertEqual(s[0].notes_slide.notes_text_frame.text, "Open with the headline.")
        body = [sh for sh in s[1].placeholders if sh.placeholder_format.idx == 1][0]
        self.assertEqual([(p.text, p.level) for p in body.text_frame.paragraphs],
                         [("Faster builds", 0), ("cache hits up", 1), ("fewer retries", 1), ("New owner", 0)])
        tbl = next(sh for sh in s[2].shapes if sh.has_table).table
        self.assertEqual((tbl.cell(0, 2).text, tbl.cell(1, 1).text, tbl.cell(2, 2).text), ("Q3", "12", ""))
        chart = next(sh for sh in s[3].shapes if sh.has_chart).chart
        self.assertEqual(chart.chart_type, XL_CHART_TYPE.LINE_MARKERS)
        self.assertEqual(list(chart.plots[0].categories), ["Jan", "Feb", "Mar"])
        self.assertFalse(chart.has_legend)
        self.assertTrue(any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in s[4].shapes))
        for i in (2, 3, 4):   # table, chart, image: content sits wholly below a top title
            title = s[i].shapes.title
            content = next(sh for sh in s[i].shapes if sh.has_table or sh.has_chart or sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
            self.assertNotIn("title slide", s[i].slide_layout.name.lower())
            self.assertLessEqual(title.top + title.height, content.top, f"slide {i + 1}: title overlaps content")
            self.assertLessEqual(content.top + content.height, prs.slide_height, f"slide {i + 1}: content runs off the slide")
        for slide in s:   # no leftover "Click to add text" boxes
            self.assertFalse([ph for ph in slide.placeholders if ph.has_text_frame and not ph.text_frame.text.strip()])

    def test_lint_and_template(self):
        _, warnings = self.bd.build(self.outline, self.d)
        self.assertTrue(any("slide 6: title longer" in w for w in warnings))
        self.assertTrue(any("slide 6: 9 bullets" in w for w in warnings))
        self.assertFalse(any(w.startswith("slide 2") for w in warnings))
        from pptx import Presentation
        tpl = Presentation()
        tpl.slides.add_slide(tpl.slide_layouts[0]).shapes.title.text = "SAMPLE SLIDE"
        tpl.save(os.path.join(self.d, "brand.pptx"))
        prs = self.build({"template": "brand.pptx", "slides": [{"layout": "title only", "title": "Only me"}]})
        self.assertEqual([sl.shapes.title.text for sl in prs.slides], ["Only me"])
        with self.assertRaises(SystemExit):
            self.bd.build({"slides": [{"layout": "No Such Layout", "title": "x"}]})
        with self.assertRaises(SystemExit):
            self.bd.build({"template": "brand.pptm", "slides": []}, self.d)


if __name__ == "__main__":
    unittest.main()

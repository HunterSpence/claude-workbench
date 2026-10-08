#!/usr/bin/env python3
"""build_deck.py: build a .pptx from a JSON outline with python-pptx, optionally on a template.

  python build_deck.py layouts <template.pptx>          list layouts and their placeholders
  python build_deck.py build <outline.json> <out.pptx>  build the deck, then print lint warnings

Outline (JSON):
{
  "template": "brand.pptx",              optional; its sample slides are removed unless "keep_template_slides": true
  "slides": [
    {"layout": "Title Slide", "title": "Q3 review", "subtitle": "Platform team", "notes": "Open with the headline."},
    {"title": "What changed", "bullets": ["Point", ["sub-point", "sub-point"], "Point"]},
    {"title": "Costs", "table": {"columns": ["Item", "Q2", "Q3"], "rows": [ ["Compute", "12", "9"] ]}},
    {"title": "Trend", "chart": {"type": "line", "categories": ["Jan", "Feb"], "series": {"Users": [10, 14]},
                                 "number_format": "0", "legend": false}},
    {"title": "Architecture", "image": "diagram.png"}
  ]
}
Layouts are matched by name (case-insensitive); without "layout" a sensible default is chosen from
the template's own layouts. Relative paths resolve against the outline's folder.
"""
import json, os, sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Pt

CHARTS = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
          "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE_MARKERS,
          "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
LINT = {"title_chars": 70, "bullets": 6, "bullet_chars": 110, "table_rows": 12}


def ph_types(layout):
    return [ph.placeholder_format.type for ph in layout.placeholders]


def pick_layout(prs, spec):
    layouts = list(prs.slide_layouts)
    if spec.get("layout"):
        for lo in layouts:
            if lo.name.lower() == spec["layout"].lower():
                return lo
        raise SystemExit(f"layout '{spec['layout']}' not in template; have: {[lo.name for lo in layouts]}")

    def has(lo, kinds):
        return any(t in kinds for t in ph_types(lo))

    # Content slides need a TOP title (PP_PLACEHOLDER.TITLE). A centred title (title-slide or section
    # layouts) sits mid-slide, so a table or chart under it would be drawn over it.
    top_title = {PP_PLACEHOLDER.TITLE}
    no_extras = lambda lo: not has(lo, BODY_TYPES | {PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.CENTER_TITLE})
    if spec.get("subtitle") is not None:
        prefs = [lambda lo: PP_PLACEHOLDER.CENTER_TITLE in ph_types(lo) or PP_PLACEHOLDER.SUBTITLE in ph_types(lo)]
    elif spec.get("bullets") is not None:
        prefs = [lambda lo: has(lo, top_title) and has(lo, BODY_TYPES)]
    else:   # table, chart, image or title only: top title and nothing else, else title + content
        prefs = [lambda lo: has(lo, top_title) and no_extras(lo), lambda lo: has(lo, top_title) and has(lo, BODY_TYPES)]
    for want in prefs + [lambda lo: has(lo, top_title), lambda lo: has(lo, TITLE_TYPES)]:
        for lo in layouts:
            if want(lo):
                return lo
    return layouts[0]


def drop_all_slides(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def content_box(prs, slide):
    """Geometry for a table/chart/image: the layout's body area if it has one, else the space
    below the title (never overlapping it), inside side and bottom margins."""
    for ph in slide.placeholders:
        if ph.placeholder_format.type in BODY_TYPES:
            return ph.left, ph.top, ph.width, ph.height
    W, H = prs.slide_width, prs.slide_height
    top = int(H * 0.06)
    for ph in slide.placeholders:
        if ph.placeholder_format.type in TITLE_TYPES:
            top = max(top, ph.top + ph.height)
    top += int(H * 0.03)
    bottom = int(H * 0.95)
    return Emu(int(W * 0.06)), Emu(top), Emu(int(W * 0.88)), Emu(max(bottom - top, int(H * 0.2)))


def fill_bullets(tf, items, level=0):
    for item in items:
        if isinstance(item, list):
            fill_bullets(tf, item, level + 1)
            continue
        p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].text) else tf.add_paragraph()
        p.text, p.level = str(item), min(level, 4)


def add_table(slide, box, spec):
    cols, rows = spec["columns"], spec.get("rows", [])
    x, y, w, h = box
    row_h = min(int(h / (len(rows) + 1)), int(Pt(32)))
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), x, y, w, Emu(row_h * (len(rows) + 1)))
    tbl = shape.table
    for c, name in enumerate(cols):
        tbl.cell(0, c).text = str(name)
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = "" if val is None else str(val)
    return shape


def add_chart(slide, box, spec):
    data = CategoryChartData()
    data.categories = spec["categories"]
    for name, values in spec["series"].items():
        data.add_series(name, values, number_format=spec.get("number_format"))
    kind = CHARTS[spec.get("type", "column")]
    chart = slide.shapes.add_chart(kind, *box, data).chart
    many = len(spec["series"]) > 1 or kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    chart.has_legend = spec.get("legend", many)
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    if spec.get("data_labels"):
        plot = chart.plots[0]
        plot.has_data_labels = True
        if spec.get("number_format"):
            plot.data_labels.number_format, plot.data_labels.number_format_is_linked = spec["number_format"], False
    chart.has_title = bool(spec.get("chart_title"))   # the slide title carries the message by default
    if chart.has_title:
        chart.chart_title.text_frame.text = spec["chart_title"]
    return chart


def add_image(slide, box, path):
    from PIL import Image   # Pillow ships with python-pptx
    x, y, w, h = box
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, Emu(x + (w - pw) // 2), Emu(y + (h - ph) // 2), Emu(pw), Emu(ph))


def remove_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph.element.getparent().remove(ph.element)


def lint(spec, n):
    w = []
    if len(spec.get("title", "")) > LINT["title_chars"]:
        w.append(f"slide {n}: title longer than {LINT['title_chars']} chars; it may wrap or overflow")
    flat = []
    stack = list(spec.get("bullets") or [])
    while stack:
        b = stack.pop(0)
        stack[:0] = b if isinstance(b, list) else []
        if not isinstance(b, list):
            flat.append(str(b))
    if len(flat) > LINT["bullets"]:
        w.append(f"slide {n}: {len(flat)} bullets (> {LINT['bullets']}); split the slide")
    w += [f"slide {n}: bullet over {LINT['bullet_chars']} chars: '{b[:40]}...'" for b in flat if len(b) > LINT["bullet_chars"]]
    if spec.get("table") and len(spec["table"].get("rows", [])) > LINT["table_rows"]:
        w.append(f"slide {n}: table has more than {LINT['table_rows']} rows; text will be tiny")
    if not spec.get("title"):
        w.append(f"slide {n}: no title (bad for navigation and accessibility)")
    return w


def build(outline, base_dir="."):
    path = lambda p: p if os.path.isabs(p) else os.path.join(base_dir, p)
    if outline.get("template") and not str(outline["template"]).lower().endswith(".pptx"):
        raise SystemExit("template must be a .pptx: save a .potx/.pptm as .pptx in PowerPoint first "
                         "(macro-enabled and legacy files are refused)")
    prs = Presentation(path(outline["template"])) if outline.get("template") else Presentation()
    if outline.get("template") and not outline.get("keep_template_slides"):
        drop_all_slides(prs)
    warnings = []
    for n, spec in enumerate(outline["slides"], 1):
        slide = prs.slides.add_slide(pick_layout(prs, spec))
        title_ph = next((ph for ph in slide.placeholders if ph.placeholder_format.type in TITLE_TYPES), None)
        if spec.get("title") and title_ph is not None:
            title_ph.text_frame.text = spec["title"]
        elif spec.get("title"):
            warnings.append(f"slide {n}: layout has no title placeholder; title dropped")
        if spec.get("subtitle") is not None:
            sub = next((ph for ph in slide.placeholders if ph.placeholder_format.type in (PP_PLACEHOLDER.SUBTITLE, *BODY_TYPES)), None)
            if sub is not None:
                sub.text_frame.text = spec["subtitle"]
        if spec.get("bullets") is not None:
            body = next((ph for ph in slide.placeholders if ph.placeholder_format.type in BODY_TYPES), None)
            if body is None:
                raise SystemExit(f"slide {n}: layout '{slide.slide_layout.name}' has no body placeholder for bullets")
            fill_bullets(body.text_frame, spec["bullets"])
        box = content_box(prs, slide)
        if spec.get("table"):
            add_table(slide, box, spec["table"])
        if spec.get("chart"):
            add_chart(slide, box, spec["chart"])
        if spec.get("image"):
            add_image(slide, box, path(spec["image"]))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        remove_empty_placeholders(slide)
        warnings += lint(spec, n)
    return prs, warnings


def main(argv=None):
    a = argv if argv is not None else sys.argv[1:]
    if len(a) == 2 and a[0] == "layouts":
        for i, lo in enumerate(Presentation(a[1]).slide_layouts):
            phs = ", ".join(f"{ph.placeholder_format.idx}:{ph.placeholder_format.type}" for ph in lo.placeholders)
            print(f"[{i}] {lo.name}  ->  {phs}")
        return 0
    if len(a) == 3 and a[0] == "build":
        with open(a[1], encoding="utf-8") as f:
            outline = json.load(f)
        prs, warnings = build(outline, os.path.dirname(os.path.abspath(a[1])))
        prs.save(a[2])
        print(f"wrote {a[2]} ({len(prs.slides)} slides)")
        for w in warnings:
            print("LINT " + w)
        print("Next: render it (office-render skill) and look at every slide.")
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main())

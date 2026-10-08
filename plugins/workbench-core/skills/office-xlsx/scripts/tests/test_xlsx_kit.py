import importlib.util, os, shutil, sys, tempfile, unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
COM = os.name == "nt" and os.environ.get("OFFICE_COM_TESTS") == "1"   # opt-in: launches real Excel


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "needs openpyxl")
class XlsxKitTest(unittest.TestCase):
    def setUp(self):
        import xlsx_kit
        from openpyxl import Workbook
        self.k = xlsx_kit
        self.d = tempfile.mkdtemp()
        self.path = os.path.join(self.d, "model.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Model"
        ws.append(["Item", "Units", "Price", "Revenue"])
        for i, (name, u, p) in enumerate([("A", 10, 2.5), ("B", 4, 10), ("C", 7, 3), ("D", 1, 1)], start=2):
            ws.append([name, u, p, f"=B{i}*C{i}"])
        ws["D4"] = 40                                  # a typed result where a formula belongs
        ws.append(["Total", "=SUM(B2:B5)", None, "=SUM(D2:D5)"])
        ws["F1"], ws["F2"] = "Ratio", "=D6/0"
        ws2 = wb.create_sheet("Summary")
        ws2.append(["Metric", "Value"])
        ws2.append(["Revenue", "=Model!D6"])
        ws2["C2"] = "#DIV/0!"                          # openpyxl stores this as an error value
        self.k.add_list_validation(ws2, "D2:D10", ["Low", "Medium", "High"], prompt="Pick a risk level")
        self.k.style_workbook(wb)
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def test_styles_freeze_validation(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        m, s = wb["Model"], wb["Summary"]
        self.assertEqual(m["B2"].style, "Input")
        self.assertEqual(m["B2"].font.color.rgb, "000000FF")
        self.assertEqual(m["D2"].style, "Formula")
        self.assertEqual(m["A1"].style, "Header")
        self.assertEqual(m["D6"].style, "Total")
        self.assertEqual(s["B2"].style, "Link")
        self.assertEqual(s["B2"].font.color.rgb, "00008000")
        self.assertEqual(m.freeze_panes, "A2")
        self.assertEqual(s.data_validations.dataValidation[0].formula1, '"Low,Medium,High"')

    def test_check_without_recalc(self):
        problems = self.k.check(self.path)
        kinds = {(p["kind"], p["sheet"], p["cell"]) for p in problems}
        self.assertIn(("error", "Summary", "C2"), kinds)
        self.assertIn(("uncalculated", "Model", "D2"), kinds)
        self.assertIn(("hardcoded", "Model", "D4"), kinds)
        self.assertNotIn(("hardcoded", "Model", "B3"), kinds)       # an input column, no formulas around it
        self.assertEqual(self.k.main(["check", self.path]), 1)

    def test_refuses_macro_and_legacy_formats(self):
        for ext in (".xlsm", ".xltm", ".xls"):
            p = os.path.join(self.d, "book" + ext)
            shutil.copyfile(self.path, p)
            self.assertEqual(self.k.main(["check", p, "--recalc", "excel"]), 2)
            self.assertEqual(self.k.main(["style", p]), 2)
            with self.assertRaises(ValueError):
                self.k.recalc(p, "excel")
        ps = self.k.PS_RECALC
        self.assertLess(ps.index("AutomationSecurity = 3"), ps.index(".Open("))
        self.assertIn("Workbooks.Open($env:XLSX_IN, 0, $true)", ps)       # read-only

    def test_xlfn_prefix(self):
        f = self.k.prefix_formula
        self.assertEqual(f("=XLOOKUP(A2,B:B,C:C)"), ("=_xlfn.XLOOKUP(A2,B:B,C:C)", None))
        self.assertEqual(f("=filter(A2:A9, B2:B9>1)")[0], "=_xlfn._xlws.FILTER(A2:A9, B2:B9>1)")
        self.assertEqual(f("=SORTBY(A2:A9,B2:B9)")[0], "=_xlfn.SORTBY(A2:A9,B2:B9)")
        self.assertEqual(f("=STDEV.S(A1:A5)+COUNTIFS(A:A,1)")[0], "=_xlfn.STDEV.S(A1:A5)+COUNTIFS(A:A,1)")
        for same in ('="XLOOKUP(" & A1', "=_xlfn.XLOOKUP(A2,B:B,C:C)", "=SUM(A1:A3)", "plain text", 42):
            self.assertEqual(f(same), (same, None))
        new, warn = f("=LET(x, A1, x*2)")
        self.assertEqual(new, "=LET(x, A1, x*2)")
        self.assertIn("_xlpm.", warn)
        from openpyxl import load_workbook
        wb = load_workbook(self.path)
        wb["Model"]["G2"] = "=XLOOKUP(A2,A2:A5,D2:D5)"
        wb["Model"]["G3"] = "=LAMBDA(x,x+1)(2)"
        wb.save(self.path)
        self.assertIn(("xlfn", "Model", "G2"), {(p["kind"], p["sheet"], p["cell"]) for p in self.k.check(self.path)})
        warnings = self.k.prefix_new_functions(wb)
        self.assertEqual(wb["Model"]["G2"].value, "=_xlfn.XLOOKUP(A2,A2:A5,D2:D5)")
        self.assertEqual([c for c, _ in warnings], ["Model!G3"])

    def _inject(self, path, parts):
        import zipfile
        with zipfile.ZipFile(path, "a") as z:
            for name, data in parts.items():
                z.writestr(name, data)

    def test_roundtrip_loss_warns_and_style_refuses(self):
        self.assertEqual(self.k.roundtrip_loss(self.path), [])
        self._inject(self.path, {
            "xl/drawings/drawing9.xml": '<xdr:wsDr xmlns:xdr="x"><xdr:twoCellAnchor><xdr:sp><x/></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>',
            "xl/threadedComments/threadedComment1.xml": "<ThreadedComments/>",
            "xl/pivotTables/pivotTable1.xml": "<pivotTableDefinition/>"})
        lost = {k for k, _, _ in self.k.roundtrip_loss(self.path)}
        self.assertTrue({"drawings", "shapes", "threaded comments", "pivot tables"} <= lost, lost)
        before = open(self.path, "rb").read()
        self.assertEqual(self.k.main(["style", self.path]), 3)
        self.assertEqual(open(self.path, "rb").read(), before, "nothing written when parts would be lost")
        self.assertEqual(self.k.main(["style", self.path, "--accept-loss"]), 3, "--accept-loss needs --out")
        out = os.path.join(self.d, "styled.xlsx")
        self.assertEqual(self.k.main(["style", self.path, "--accept-loss", "--out", out]), 0)
        self.assertTrue(os.path.exists(out))
        self.assertEqual(self.k.main(["risk", self.path]), 3)

    def test_openpyxl_charts_comments_images_survive(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        wb = Workbook()
        ws = wb.active
        for r in (["k", "v"], ["a", 1], ["b", 2]):
            ws.append(r)
        ws["B2"].comment = Comment("note", "me")
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        ws.add_chart(ch, "D2")
        if importlib.util.find_spec("PIL"):
            from PIL import Image as PILImage
            from openpyxl.drawing.image import Image
            png = os.path.join(self.d, "p.png")
            PILImage.new("RGB", (20, 20)).save(png)
            ws.add_image(Image(png), "D20")
        p = os.path.join(self.d, "kept.xlsx")
        wb.save(p)
        inv = self.k.inventory(p)
        self.assertEqual((inv["charts"], inv["comments"]), (1, 1))
        self.assertEqual(self.k.roundtrip_loss(p), [], "no false alarm on parts openpyxl keeps")

    @unittest.skipUnless(COM, "set OFFICE_COM_TESTS=1 on Windows with Excel")
    def test_excel_recalc_finds_real_errors(self):
        self.assertEqual(self.k.recalc(self.path, "excel"), "excel")
        problems = self.k.check(self.path)
        self.assertIn(("error", "Model", "F2"), {(p["kind"], p["sheet"], p["cell"]) for p in problems})
        self.assertFalse([p for p in problems if p["kind"] == "uncalculated"])
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(self.path, data_only=True)["Model"]["D2"].value, 25)


if __name__ == "__main__":
    unittest.main()

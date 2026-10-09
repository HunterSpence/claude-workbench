import os

REQUIRE = os.environ.get("WB_REQUIRE_ALL") == "1"   # house convention: a missing dependency fails instead of skipping

import glob  # noqa: E402
import json  # noqa: E402
import sys  # noqa: E402
import tempfile  # noqa: E402
import unittest  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TEMPLATES = sorted(glob.glob(os.path.join(os.path.dirname(os.path.dirname(HERE)), "templates", "*.json")))


class Templates(unittest.TestCase):
    def test_starter_specs_build_with_dropdowns_and_live_formulas(self):
        import xlsx_kit
        from openpyxl import load_workbook
        self.assertEqual([os.path.basename(t) for t in TEMPLATES], ["raid-log.json", "traceability-matrix.json"])
        with tempfile.TemporaryDirectory() as d:
            for t in TEMPLATES:
                out = os.path.join(d, os.path.basename(t) + ".xlsx")
                self.assertEqual(xlsx_kit.main(["from-spec", t, out]), 0, t)
                wb = load_workbook(out)
                spec = json.load(open(t, encoding="utf-8"))
                self.assertEqual(wb.sheetnames, [s["name"] for s in spec["sheets"]])
                self.assertTrue(wb.worksheets[0].data_validations.dataValidation, t)
                self.assertTrue(any(str(c.value).startswith("=COUNT") for row in wb.worksheets[-1].iter_rows() for c in row), t)
            self.assertEqual(xlsx_kit.main(["from-spec", TEMPLATES[0], os.path.join(d, "x.csv")]), 2)


if __name__ == "__main__":
    unittest.main()

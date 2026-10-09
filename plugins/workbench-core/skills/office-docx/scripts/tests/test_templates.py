import os

REQUIRE = os.environ.get("WB_REQUIRE_ALL") == "1"   # house convention: a missing dependency fails instead of skipping

import glob  # noqa: E402
import sys  # noqa: E402
import tempfile  # noqa: E402
import unittest  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TEMPLATES = sorted(glob.glob(os.path.join(os.path.dirname(os.path.dirname(HERE)), "templates", "*.json")))


class Templates(unittest.TestCase):
    def test_starter_outlines_build_with_tables_and_a_provenance_section(self):
        import build_doc
        from docx import Document
        self.assertEqual([os.path.basename(t) for t in TEMPLATES], ["decision-log.json", "requirements-pack.json"])
        with tempfile.TemporaryDirectory() as d:
            for t in TEMPLATES:
                out = os.path.join(d, os.path.basename(t) + ".docx")
                self.assertEqual(build_doc.main(["build", t, out]), 0, t)
                doc = Document(out)
                self.assertTrue(doc.tables, t)
                self.assertIn("Provenance", [p.text for p in doc.paragraphs], t)


if __name__ == "__main__":
    unittest.main()

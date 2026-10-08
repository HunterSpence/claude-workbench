import importlib.util, json, os, shutil, sys, tempfile, unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
HAVE_DOCX = importlib.util.find_spec("docx") is not None


@unittest.skipUnless(HAVE_DOCX, "needs python-docx")
class BuildDocTest(unittest.TestCase):
    def setUp(self):
        import build_doc
        self.bd = build_doc
        self.d = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def path(self, name):
        return os.path.join(self.d, name)

    def build(self, outline):
        p = self.path("outline.json")
        with open(p, "w", encoding="utf-8") as f:
            json.dump(outline, f)
        out = self.path("out.docx")
        self.assertEqual(self.bd.main(["build", p, out]), 0)
        from docx import Document
        return Document(out)

    def png(self, w=3000, h=600):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("needs Pillow")
        Image.new("RGB", (w, h), (30, 90, 160)).save(self.path("pic.png"))
        return "pic.png"

    def test_every_block_kind(self):
        from docx.oxml.ns import qn
        doc = self.build({"header": "Team memo", "footer": "Internal", "page_numbers": True, "blocks": [
            {"title": "Release plan"},
            {"heading": "Summary", "level": 1},
            {"paragraph": "Plain **bold** and *italic* and `code` and [link](https://example.com)."},
            {"bullets": ["one", ["one-a", "one-b"], "two"]},
            {"numbered": ["first", "second"]},
            {"heading": "Detail", "level": 2},
            {"numbered": ["again one", "again two"]},
            {"table": {"columns": ["Item", "Owner"], "rows": [ ["Build", "<name>"], ["Test", None] ]}},
            {"image": self.png(), "caption": "Figure 1"},
            {"quote": "quoted"},
            {"page_break": True},
        ]})
        styles = [p.style.name for p in doc.paragraphs]
        for want in ("Title", "Heading 1", "Heading 2", "List Bullet", "List Bullet 2", "List Number", "Caption", "Quote"):
            self.assertIn(want, styles)
        para = next(p for p in doc.paragraphs if p.text.startswith("Plain"))
        self.assertTrue(any(r.bold and r.text == "bold" for r in para.runs))
        self.assertTrue(any(r.italic and r.text == "italic" for r in para.runs))
        self.assertTrue(para._p.findall(qn("w:hyperlink")), "markdown link becomes a real hyperlink")
        self.assertIn("link", para.text)
        # each numbered list restarts: two different numIds, each with startOverride 1
        num_of = lambda p: p._p.pPr.numPr.numId.val if p._p.pPr is not None and p._p.pPr.numPr is not None else "style"
        nums = {num_of(p) for p in doc.paragraphs if p.style.name == "List Number"}
        self.assertEqual(len(nums), 2)
        t = doc.tables[0]
        self.assertTrue(all(r.bold for r in t.rows[0].cells[0].paragraphs[0].runs))
        self.assertTrue(t.rows[0]._tr.trPr.findall(qn("w:tblHeader")), "header row repeats")
        self.assertEqual(t.rows[2].cells[1].text, "")
        sec = doc.sections[0]
        self.assertLessEqual(doc.inline_shapes[0].width, sec.page_width - sec.left_margin - sec.right_margin,
                             "a wide image is scaled to the page")
        self.assertEqual(sec.header.paragraphs[0].text, "Team memo")
        self.assertTrue(sec.footer.paragraphs[0]._p.findall(qn("w:fldSimple")), "page number field present")

    def test_markdown_and_lint(self):
        md = self.path("in.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("# Memo\n\nIntro with **bold**.\n\n### Skipped a level\n\n- a\n    - a1\n- b\n\n"
                    "1. x\n2. y\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n![missing](nope.png)\n")
        out = self.path("md.docx")
        from io import StringIO
        from contextlib import redirect_stdout
        buf = StringIO()
        with redirect_stdout(buf):
            self.assertEqual(self.bd.main(["md", md, out]), 0)
        printed = buf.getvalue()
        self.assertIn("LINT heading 'Skipped a level' jumps from level 1 to 3", printed)
        self.assertIn("LINT image not found", printed)
        from docx import Document
        doc = Document(out)
        styles = [p.style.name for p in doc.paragraphs]
        self.assertEqual(styles.count("List Bullet"), 2)
        self.assertEqual(styles.count("List Bullet 2"), 1)
        self.assertEqual(len(doc.tables), 1)

    def test_template_cleared_and_styles(self):
        from docx import Document
        tpl = Document()
        tpl.add_paragraph("SAMPLE CONTENT")
        tpl.sections[0].header.paragraphs[0].text = "Brand header"
        tp = self.path("tpl.docx")
        tpl.save(tp)
        doc = self.build({"template": "tpl.docx", "blocks": [{"heading": "Hello"}]})
        self.assertNotIn("SAMPLE CONTENT", [p.text for p in doc.paragraphs])
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "Brand header", "template header kept")
        from io import StringIO
        from contextlib import redirect_stdout
        buf = StringIO()
        with redirect_stdout(buf):
            self.bd.main(["styles", tp])
        self.assertIn("Heading 1", buf.getvalue())

    def test_refusals(self):
        for bad in ("t.docm", "t.dotm", "t.doc", "t.rtf"):
            with self.assertRaises(SystemExit):
                self.bd.Builder(self.path(bad), base_dir=self.d)
        with self.assertRaises(SystemExit):
            self.bd.Builder(self.path("t.dotx"), base_dir=self.d)
        with self.assertRaises(SystemExit):
            self.bd.to_markdown(self.path("x.docm"))
        b = self.bd.Builder()
        with self.assertRaises(SystemExit):
            b.save(self.path("out.docm"))

    def test_paths_stay_inside_the_outline_folder(self):
        os.makedirs(self.path("img"))
        self.png()
        shutil.move(self.path("pic.png"), self.path("img/pic.png"))
        self.assertEqual(len(self.build({"blocks": [{"image": "img/pic.png"}]}).inline_shapes), 1, "a subfolder is fine")
        outside = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, outside, True)
        shutil.copy(self.path("img/pic.png"), os.path.join(outside, "x.png"))
        rel_out = os.path.relpath(os.path.join(outside, "x.png"), self.d)
        p = self.path("outline.json")
        for bad in (rel_out, os.path.join(outside, "x.png"), "../x.png", "\\\\server\\share\\x.png", "//server/share/x.png",
                    "http://example.com/x.png", "file:///C:/x.png"):
            with open(p, "w", encoding="utf-8") as f:
                json.dump({"blocks": [{"image": bad}]}, f)
            with self.assertRaises(SystemExit, msg=bad) as cm:
                self.bd.main(["build", p, self.path("out.docx")])
            self.assertIn("refusing image", str(cm.exception), bad)
        for tpl in (os.path.join(outside, "t.docx"), "\\\\server\\share\\t.docx"):
            with open(p, "w", encoding="utf-8") as f:
                json.dump({"template": tpl, "blocks": []}, f)
            with self.assertRaises(SystemExit) as cm:
                self.bd.main(["build", p, self.path("out.docx")])
            self.assertIn("refusing template", str(cm.exception))
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"blocks": [{"image": os.path.join(outside, "x.png")}]}, f)
        self.assertEqual(self.bd.main(["build", p, self.path("out.docx"), "--allow-outside"]), 0)
        md = self.path("doc.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("# T\n\n![cap](../x.png)\n")
        with self.assertRaises(SystemExit):
            self.bd.main(["md", md, self.path("o2.docx")])

    @unittest.skipUnless(importlib.util.find_spec("markitdown"), "needs markitdown")
    def test_round_trip_to_markdown(self):
        self.build({"blocks": [{"heading": "Round trip"}, {"bullets": ["alpha", "beta"]},
                               {"table": {"columns": ["K", "V"], "rows": [ ["a", "1"] ]}}]})
        md = self.bd.to_markdown(self.path("out.docx"))
        self.assertIn("Round trip", md)
        self.assertIn("alpha", md)
        self.assertIn("| a | 1 |", md)


class NestTest(unittest.TestCase):
    def test_nest(self):
        if not HAVE_DOCX:
            self.skipTest("needs python-docx")
        import build_doc
        self.assertEqual(build_doc.nest([(0, "a"), (2, "a1"), (2, "a2"), (4, "x"), (0, "b"), (2, "b1")]),
                         ["a", ["a1", "a2", ["x"]], "b", ["b1"]])
        self.assertEqual(build_doc.nest([(0, "a"), (8, "deep"), (0, "b")]), ["a", ["deep"], "b"])


if __name__ == "__main__":
    unittest.main()

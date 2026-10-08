import importlib.util, os, shutil, sys, tempfile, unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import render

HAVE_PDFIUM = importlib.util.find_spec("pypdfium2") is not None and importlib.util.find_spec("PIL") is not None
COM = os.name == "nt" and os.environ.get("OFFICE_COM_TESTS") == "1"   # opt-in: launches real Office apps


def tiny_pdf(path, pages=2):
    """A valid PDF with `pages` letter-size pages, each holding a filled rectangle."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            "<< /Type /Pages /Kids [" + " ".join(f"{3 + 2 * i} 0 R" for i in range(pages)) + f"] /Count {pages} >>"]
    for i in range(pages):
        stream = f"0 0 1 rg {50 + 40 * i} 500 200 150 re f"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)


class RenderTest(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    @unittest.skipUnless(HAVE_PDFIUM, "needs pypdfium2 and Pillow")
    def test_pdf_to_png_and_stale_cleanup(self):
        from PIL import Image
        src = os.path.join(self.d, "doc.pdf")
        tiny_pdf(src, 2)
        out = os.path.join(self.d, "out")
        os.makedirs(out)
        open(os.path.join(out, "doc-09.png"), "wb").close()      # a stale render from an older, longer version
        paths, used = render.render(src, out, dpi=72)
        self.assertEqual(used, "pdf")
        self.assertEqual([os.path.basename(p) for p in paths], ["doc-01.png", "doc-02.png"])
        self.assertFalse(os.path.exists(os.path.join(out, "doc-09.png")))
        with Image.open(paths[0]) as im:
            self.assertEqual(im.size, (612, 792))
            self.assertEqual(im.getpixel((150, 792 - 575))[:3], (0, 0, 255))   # inside the blue rectangle
            self.assertEqual(im.getpixel((10, 10))[:3], (255, 255, 255))
        self.assertEqual(len(render.render(src, out, dpi=36, first=1)[0]), 1)

    def test_refuses_macro_and_legacy_formats(self):
        for ext in (".xlsm", ".xltm", ".docm", ".dotm", ".pptm", ".potm", ".doc", ".xls", ".ppt", ".rtf"):
            src = os.path.join(self.d, "f" + ext)
            open(src, "wb").close()
            with self.assertRaisesRegex(ValueError, "macro-enabled or legacy"):
                render.render(src, os.path.join(self.d, "out"))
            self.assertEqual(render.main([src, "--out", self.d]), 1)
        self.assertIn("AutomationSecurity = 3", render.PS["word"])
        for app in render.PS.values():   # macros off before any Open, every app
            self.assertLess(app.index("AutomationSecurity = 3"), app.index(".Open("))

    @unittest.skipIf(render.find_soffice(), "LibreOffice is installed")
    def test_missing_engine_fails_loudly(self):
        src = os.path.join(self.d, "x.pptx")
        open(src, "wb").close()
        with self.assertRaisesRegex(RuntimeError, "not found"):
            render.to_pdf(src, os.path.join(self.d, "x.pdf"), "libreoffice")
        self.assertEqual(render.main([src, "--out", self.d, "--engine", "libreoffice"]), 1)

    @unittest.skipUnless(COM and HAVE_PDFIUM and importlib.util.find_spec("pptx"), "set OFFICE_COM_TESTS=1 on Windows with Office")
    def test_office_com_pptx_docx_xlsx(self):
        import docx, openpyxl, pptx
        p = pptx.Presentation()
        for t in ("One", "Two"):
            p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = t
        p.save(os.path.join(self.d, "deck.pptx"))
        d = docx.Document()
        d.add_paragraph("Hello")
        d.save(os.path.join(self.d, "memo.docx"))
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        wb.save(os.path.join(self.d, "book.xlsx"))
        for name, pages in (("deck.pptx", 2), ("memo.docx", 1), ("book.xlsx", 1)):
            paths, used = render.render(os.path.join(self.d, name), os.path.join(self.d, "out"), engine="office", dpi=50)
            self.assertEqual((used, len(paths)), ("office", pages), name)


if __name__ == "__main__":
    unittest.main()

import os

REQUIRE = os.environ.get("WB_REQUIRE_ALL") == "1"   # house convention: a missing dependency fails instead of skipping

import hashlib  # noqa: E402
import json  # noqa: E402
import shutil  # noqa: E402
import subprocess  # noqa: E402
import sys  # noqa: E402
import tempfile  # noqa: E402
import unittest  # noqa: E402
import zipfile  # noqa: E402
from pathlib import Path  # noqa: E402

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))


def need(module):
    try:
        return __import__(module)
    except ImportError:
        if REQUIRE:
            raise
        return None


etree = need("lxml.etree") and sys.modules["lxml.etree"]
if etree is None:
    raise unittest.SkipTest("needs lxml")
docx, openpyxl = need("docx"), need("openpyxl")
import render  # noqa: E402
import scrub  # noqa: E402

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
RNS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
OD = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
PK = "http://schemas.openxmlformats.org/package/2006/relationships/"
MS = "http://schemas.microsoft.com/office/"
XML = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
CTYPE = "application/vnd.openxmlformats-officedocument."
TOKENS = ["Jordan Example", "Casey Example", "Example Holdings", "Robin Example", "REMARK-IN-PROPERTIES", "PROJECT-CODE-VALUE"]


def rels(*items):
    body = "".join(f'<Relationship Id="{i}" Type="{t}" Target="{g}"/>' for i, t, g in items)
    return f'{XML}<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{body}</Relationships>'


CORE = (f'{XML}<cp:coreProperties xmlns:cp="{PK[:-14]}metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>Plan</dc:title><dc:creator>Jordan Example</dc:creator><cp:lastModifiedBy>Casey Example</cp:lastModifiedBy>"
        "<dc:description>REMARK-IN-PROPERTIES</dc:description></cp:coreProperties>")
VT = 'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"'
APP = (f'{XML}<Properties xmlns="{OD[:-14]}extended-properties" {VT}><Application>Test</Application>'
       "<Company>Example Holdings</Company><Manager>Robin Example</Manager>{extra}</Properties>")
CUSTOM = (f'{XML}<Properties xmlns="{OD[:-14]}custom-properties" {VT}>'
          '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="MSIP_Label_0001_Name"><vt:lpwstr>Internal</vt:lpwstr></property>'
          '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="ProjectCode"><vt:lpwstr>PROJECT-CODE-VALUE</vt:lpwstr></property>'
          "</Properties>")
LABEL = f'{XML}<clbl:labelList xmlns:clbl="{MS}2020/mipLabelMetadata"><clbl:label id="{{0001}}" enabled="1" method="Standard"/></clbl:labelList>'
PNG = b"\x89PNG\r\n\x1a\n" + b"\0" * 16


def make(path, main, parts, types, app_extra=""):
    """Write a package: `parts` name -> text/bytes, `types` name -> content type for Override entries."""
    common = {"_rels/.rels": rels(("r1", OD + "officeDocument", main), ("r2", PK + "metadata/core-properties", "docProps/core.xml"),
                                  ("r3", OD + "extended-properties", "docProps/app.xml"), ("r4", OD + "custom-properties", "docProps/custom.xml"),
                                  ("r5", PK + "metadata/thumbnail", "docProps/thumbnail.png"),
                                  ("r6", MS + "2020/02/relationships/classificationlabels", "clientMetaData/LabelInfo.xml")),
              "docProps/core.xml": CORE, "docProps/app.xml": APP.replace("{extra}", app_extra), "docProps/custom.xml": CUSTOM,
              "docProps/thumbnail.png": PNG, "clientMetaData/LabelInfo.xml": LABEL}
    ct_types = {"/docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
                "/docProps/app.xml": CTYPE + "extended-properties+xml", "/docProps/custom.xml": CTYPE + "custom-properties+xml",
                "/clientMetaData/LabelInfo.xml": "application/vnd.ms-office.classificationlabels+xml"}
    ct_types.update({"/" + k: v for k, v in types.items()})
    ct = (f'{XML}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
          '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
          + "".join(f'<Override PartName="{k}" ContentType="{v}"/>' for k, v in ct_types.items()) + "</Types>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        for name, data in {**common, **parts}.items():
            z.writestr(name, data)


def make_docx(path):
    body = (
        "<w:p><w:r><w:t>KEEP-ME paragraph.</w:t></w:r></w:p>"
        '<w:p><w:commentRangeStart w:id="0"/><w:r><w:t>Commented text stays.</w:t></w:r><w:commentRangeEnd w:id="0"/>'
        '<w:r><w:commentReference w:id="0"/></w:r></w:p>'
        '<w:p><w:ins w:id="1" w:author="Jordan Example"><w:r><w:t>INSERTED-TEXT</w:t></w:r></w:ins>'
        '<w:del w:id="2" w:author="Jordan Example"><w:r><w:delText>DELETED-SECRET</w:delText></w:r></w:del>'
        '<w:r><w:rPr><w:b/><w:rPrChange w:id="3" w:author="Casey Example"><w:rPr/></w:rPrChange></w:rPr><w:t>formatted</w:t></w:r></w:p>'
        '<w:p><w:pPr><w:rPr><w:ins w:id="4" w:author="Casey Example"/></w:rPr></w:pPr><w:r><w:t>mark</w:t></w:r></w:p>'
        "<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid>"
        '<w:tr><w:trPr><w:del w:id="5" w:author="Casey Example"/></w:trPr><w:tc><w:p><w:r><w:delText>DELETED-ROW</w:delText></w:r></w:p></w:tc></w:tr>'
        "<w:tr><w:tc><w:p><w:r><w:t>KEPT-ROW</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:sectPr/>")
    doc = (f'{XML}<w:document {W} {RNS} xmlns:w14="{MS}word/2010/wordml" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           f'mc:Ignorable="w14"><w:body>{body}</w:body></w:document>')
    comments = (f'{XML}<w:comments {W}><w:comment w:id="0" w:author="Jordan Example" w:initials="JE"><w:p><w:r><w:t>SECRET-COMMENT</w:t></w:r></w:p>'
                "</w:comment></w:comments>")
    make(path, "word/document.xml", {
        "word/document.xml": doc, "word/comments.xml": comments,
        "word/commentsExtended.xml": f'{XML}<w15:commentsEx xmlns:w15="{MS}word/2012/wordml"/>',
        "word/people.xml": f'{XML}<w15:people xmlns:w15="{MS}word/2012/wordml"><w15:person w15:author="Jordan Example"/></w15:people>',
        "word/settings.xml": f"{XML}<w:settings {W}><w:trackRevisions/></w:settings>",
        "word/styles.xml": f'{XML}<w:styles {W}>\n\n   <w:docDefaults/>\n</w:styles>',      # untouched parts keep their exact bytes
        "word/_rels/document.xml.rels": rels(("a1", OD + "comments", "comments.xml"), ("a2", MS + "2011/relationships/commentsExtended", "commentsExtended.xml"),
                                             ("a3", MS + "2011/relationships/people", "people.xml"), ("a4", OD + "settings", "settings.xml"),
                                             ("a5", OD + "styles", "styles.xml"))},
         {"word/document.xml": CTYPE + "wordprocessingml.document.main+xml", "word/comments.xml": CTYPE + "wordprocessingml.comments+xml",
          "word/commentsExtended.xml": CTYPE + "wordprocessingml.commentsExtended+xml", "word/people.xml": CTYPE + "wordprocessingml.people+xml",
          "word/settings.xml": CTYPE + "wordprocessingml.settings+xml", "word/styles.xml": CTYPE + "wordprocessingml.styles+xml"})


def make_pptx(path):
    P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    def slide(text, attrs="", extra=""):
        return f'{XML}<p:sld {P} {RNS}{attrs}><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>{extra}</p:spTree></p:cSld></p:sld>'
    link = '<p:sp><p:txBody><a:p><a:r><a:rPr><a:hlinkClick r:id="s2" action="ppaction://hlinksldjump"/></a:rPr><a:t>see hidden</a:t></a:r></a:p></p:txBody></p:sp>'
    pres = (f'{XML}<p:presentation {P} {RNS}><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="m1"/></p:sldMasterIdLst>'
            '<p:sldIdLst><p:sldId id="256" r:id="s1"/><p:sldId id="257" r:id="s2"/><p:sldId id="258" r:id="s3"/></p:sldIdLst>'
            f'<p:extLst><p:ext uri="x"><p14:sectionLst xmlns:p14="{MS}powerpoint/2010/main"><p14:section name="Main" id="{{0002}}"><p14:sldIdLst>'
            '<p14:sldId id="256"/><p14:sldId id="257"/><p14:sldId id="258"/></p14:sldIdLst></p14:section></p14:sectionLst></p:ext></p:extLst></p:presentation>')
    notes = lambda t: f"{XML}<p:notes {P}><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{t}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:notes>"  # noqa: E731
    titles = ('<HeadingPairs><vt:vector size="2" baseType="variant"><vt:variant><vt:lpstr>Slide Titles</vt:lpstr></vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant></vt:vector></HeadingPairs>'
              '<TitlesOfParts><vt:vector size="1" baseType="lpstr"><vt:lpstr>HIDDEN-SLIDE-TITLE</vt:lpstr></vt:vector></TitlesOfParts><HiddenSlides>1</HiddenSlides>')
    lay = ("ppt/slideLayouts/slideLayout1.xml", ("l1", OD + "slideMaster", "../slideMasters/slideMaster1.xml"))
    make(path, "ppt/presentation.xml", {
        "ppt/presentation.xml": pres,
        "ppt/_rels/presentation.xml.rels": rels(("m1", OD + "slideMaster", "slideMasters/slideMaster1.xml"), ("s1", OD + "slide", "slides/slide1.xml"),
                                                ("s2", OD + "slide", "slides/slide2.xml"), ("s3", OD + "slide", "slides/slide3.xml"),
                                                ("c1", OD + "commentAuthors", "commentAuthors.xml")),
        "ppt/slideMasters/slideMaster1.xml": f'{XML}<p:sldMaster {P}/>',
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": rels(("l1", OD + "slideLayout", "../slideLayouts/slideLayout1.xml")),
        lay[0]: f'{XML}<p:sldLayout {P}/>', "ppt/slideLayouts/_rels/slideLayout1.xml.rels": rels(lay[1]),
        "ppt/slides/slide1.xml": slide("VISIBLE-SLIDE", extra=link),
        "ppt/slides/_rels/slide1.xml.rels": rels(("l1", OD + "slideLayout", "../slideLayouts/slideLayout1.xml"), ("n1", OD + "notesSlide", "../notesSlides/notesSlide1.xml"),
                                                 ("c1", OD + "comments", "../comments/comment1.xml"), ("s2", OD + "slide", "slide2.xml")),
        "ppt/slides/slide2.xml": slide("HIDDEN-SLIDE-TEXT", attrs=' show="0"'),
        "ppt/slides/_rels/slide2.xml.rels": rels(("l1", OD + "slideLayout", "../slideLayouts/slideLayout1.xml"), ("n1", OD + "notesSlide", "../notesSlides/notesSlide2.xml")),
        "ppt/slides/slide3.xml": slide("SECOND-VISIBLE"),
        "ppt/slides/_rels/slide3.xml.rels": rels(("l1", OD + "slideLayout", "../slideLayouts/slideLayout1.xml")),
        "ppt/notesSlides/notesSlide1.xml": notes("SPEAKER-NOTE-ONE"),
        "ppt/notesSlides/_rels/notesSlide1.xml.rels": rels(("s1", OD + "slide", "../slides/slide1.xml")),
        "ppt/notesSlides/notesSlide2.xml": notes("HIDDEN-NOTE-TWO"),
        "ppt/notesSlides/_rels/notesSlide2.xml.rels": rels(("s1", OD + "slide", "../slides/slide2.xml")),
        "ppt/commentAuthors.xml": f'{XML}<p:cmAuthorLst {P}><p:cmAuthor id="0" name="Jordan Example" initials="JE" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>',
        "ppt/comments/comment1.xml": f'{XML}<p:cmLst {P}><p:cm authorId="0" idx="1"><p:pos x="0" y="0"/><p:text>SECRET-COMMENT-PPT</p:text></p:cm></p:cmLst>'},
        {"ppt/presentation.xml": CTYPE + "presentationml.presentation.main+xml", "ppt/slideMasters/slideMaster1.xml": CTYPE + "presentationml.slideMaster+xml",
         "ppt/slideLayouts/slideLayout1.xml": CTYPE + "presentationml.slideLayout+xml", "ppt/slides/slide1.xml": CTYPE + "presentationml.slide+xml",
         "ppt/slides/slide2.xml": CTYPE + "presentationml.slide+xml", "ppt/slides/slide3.xml": CTYPE + "presentationml.slide+xml",
         "ppt/notesSlides/notesSlide1.xml": CTYPE + "presentationml.notesSlide+xml", "ppt/notesSlides/notesSlide2.xml": CTYPE + "presentationml.notesSlide+xml",
         "ppt/commentAuthors.xml": CTYPE + "presentationml.commentAuthors+xml", "ppt/comments/comment1.xml": CTYPE + "presentationml.comments+xml"},
        app_extra=titles)


def make_xlsx(path):
    S = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    strings = (f'{XML}<sst {S}><si><t>=cmd|\' /C calc\'!A0</t></si><si><t>Hello</t></si><si><r><t>@rich</t></r><r><t> run</t></r></si></sst>')
    sheet = (f'{XML}<worksheet {S} {RNS}><sheetData>'
             '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="inlineStr"><is><t>+1+1</t></is></c><c r="C1"><v>5</v></c></row>'
             '<row r="2"><c r="A2" t="s"><v>1</v></c><c r="B2" t="inlineStr"><is><t>-2+3</t></is></c><c r="C2"><f>SUM(C1:C1)</f><v>5</v></c></row>'
             '<row r="3"><c r="A3" t="s"><v>2</v></c><c r="B3" t="inlineStr"><is><t>ok</t></is></c><c r="C3" t="str"><f>"="&amp;"text"</f><v>=text</v></c></row>'
             '<row r="4"><c r="C4" t="str"><v>=bare</v></c></row></sheetData><legacyDrawing r:id="v1"/></worksheet>')
    wb = (f'{XML}<workbook {S} {RNS} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:x15ac="{MS}spreadsheetml/2010/11/ac" mc:Ignorable="x15ac">'
          '<mc:AlternateContent><mc:Choice Requires="x15"><x15ac:absPath url="D:\\Shared\\Plans\\Heron"/></mc:Choice></mc:AlternateContent>'
          '<sheets><sheet name="Data" sheetId="1" r:id="w1"/></sheets></workbook>')
    make(path, "xl/workbook.xml", {
        "xl/workbook.xml": wb, "xl/sharedStrings.xml": strings, "xl/worksheets/sheet1.xml": sheet,
        "xl/_rels/workbook.xml.rels": rels(("w1", OD + "worksheet", "worksheets/sheet1.xml"), ("w2", OD + "sharedStrings", "sharedStrings.xml"),
                                           ("w3", MS + "2017/10/relationships/person", "persons/person.xml")),
        "xl/worksheets/_rels/sheet1.xml.rels": rels(("c1", OD + "comments", "../comments1.xml"), ("v1", OD + "vmlDrawing", "../drawings/vmlDrawing1.vml"),
                                                    ("t1", MS + "2017/10/relationships/threadedComment", "../threadedComments/threadedComment1.xml")),
        "xl/comments1.xml": f'{XML}<comments {S}><authors><author>Jordan Example</author></authors><commentList><comment ref="A1" authorId="0"><text><t>SECRET-COMMENT-XLSX</t></text></comment></commentList></comments>',
        "xl/drawings/vmlDrawing1.vml": f'{XML}<xml xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="s1">VML-COMMENT-SHAPE</v:shape></xml>',
        "xl/threadedComments/threadedComment1.xml": f'{XML}<ThreadedComments xmlns="{MS}spreadsheetml/2018/threadedcomments"><threadedComment ref="A1" id="{{0003}}"><text>THREADED-SECRET</text></threadedComment></ThreadedComments>',
        "xl/persons/person.xml": f'{XML}<personList xmlns="{MS}spreadsheetml/2018/threadedcomments"><person displayName="Jordan Example" id="{{0004}}"/></personList>'},
        {"xl/workbook.xml": CTYPE + "spreadsheetml.sheet.main+xml", "xl/sharedStrings.xml": CTYPE + "spreadsheetml.sharedStrings+xml",
         "xl/worksheets/sheet1.xml": CTYPE + "spreadsheetml.worksheet+xml", "xl/comments1.xml": CTYPE + "spreadsheetml.comments+xml",
         "xl/threadedComments/threadedComment1.xml": "application/vnd.ms-excel.threadedcomments+xml", "xl/persons/person.xml": "application/vnd.ms-excel.person+xml"})


def blob(path):
    with zipfile.ZipFile(path) as z:
        return b"\n".join(z.read(n) for n in z.namelist()).decode("utf-8", "replace")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Scrub(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.d, True)

    def run_scrub(self, maker, name, *flags):
        src, out = self.d / name, self.d / ("clean-" + name)
        maker(src)
        self.assert_valid(src)                         # the fixture itself is a sound package
        before = sha(src)
        r = subprocess.run([sys.executable, str(SCRIPTS / "scrub.py"), str(src), "--out", str(out), "--json", *flags],
                           capture_output=True, text=True, timeout=120)
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        rep = json.loads(r.stdout)
        self.assertEqual(sha(src), before, "the input must not change")
        self.assert_valid(out)
        return src, out, rep

    def assert_valid(self, path):
        """Every part parses, every relationship target and override exists, every r:id in a part resolves."""
        with zipfile.ZipFile(path) as z:
            self.assertIsNone(z.testzip())
            names = set(z.namelist())
            trees = {n: etree.fromstring(z.read(n)) for n in names if n.endswith((".xml", ".rels", ".vml"))}
        for o in trees["[Content_Types].xml"]:
            if o.tag.endswith("Override"):
                self.assertIn(o.get("PartName")[1:], names, "override without a part")
        for rn in (n for n in names if n.endswith(".rels")):
            src = scrub.source_of(rn)
            self.assertTrue(src == "" or src in names, f"{rn}: relationships for a missing part")
            ids = {r.get("Id") for r in trees[rn]}
            for r in trees[rn]:
                if r.get("TargetMode") != "External":
                    self.assertIn(scrub.target(src, r), names, f"{rn}: dangling target")
            for e in trees.get(src, etree.Element("x")).iter(etree.Element):
                for k, v in e.attrib.items():
                    if k.startswith(scrub.R):
                        self.assertIn(v, ids, f"{src}: {k} {v} has no relationship")

    def assert_gone(self, text, *tokens):
        for t in tokens:
            self.assertNotIn(t, text, t)

    def test_docx(self):
        src, out, rep = self.run_scrub(make_docx, "memo.docx")
        text = blob(out)
        self.assert_gone(text, "SECRET-COMMENT", "DELETED-SECRET", "DELETED-ROW", "commentReference", "trackRevisions", "<w:ins", "<w:del", *TOKENS)
        for kept in ("KEEP-ME", "INSERTED-TEXT", "Commented text stays.", "KEPT-ROW", "MSIP_Label_0001_Name", "labelList", 'mc:Ignorable="w14"', "xmlns:w14="):
            self.assertIn(kept, text)
        names = zipfile.ZipFile(out).namelist()
        self.assert_gone(" ".join(names), "comments.xml", "people.xml", "commentsExtended", "thumbnail")
        self.assertEqual(zipfile.ZipFile(out).read("word/styles.xml"), zipfile.ZipFile(src).read("word/styles.xml"))
        self.assertEqual((rep["changes"]["tracked_changes_accepted"], rep["changes"]["comments_removed"]), (5, 1))
        self.assertEqual(rep["label"], {"properties_kept": 1, "info_part_kept": True})
        if docx:
            self.assertIn("INSERTED-TEXT", "\n".join(p.text for p in docx.Document(str(out)).paragraphs))

    def test_pptx(self):
        src, out, rep = self.run_scrub(make_pptx, "deck.pptx")
        text = blob(out)
        self.assert_gone(text, "HIDDEN-SLIDE-TEXT", "HIDDEN-SLIDE-TITLE", "HIDDEN-NOTE-TWO", "SECRET-COMMENT-PPT", "hlinksldjump", "HiddenSlides", *TOKENS)
        for kept in ("VISIBLE-SLIDE", "SECOND-VISIBLE", "SPEAKER-NOTE-ONE", "MSIP_Label_0001_Name", "LabelInfo"):
            self.assertIn(kept, text)
        pres = etree.fromstring(zipfile.ZipFile(out).read("ppt/presentation.xml"))
        self.assertEqual([e.get("id") for e in pres.iter() if isinstance(e.tag, str) and e.tag.endswith("}sldId")], ["256", "258", "256", "258"])
        self.assertEqual(rep["changes"]["hidden_slides_removed"], 1)
        self.assertEqual(rep["changes"]["comments_removed"], 1)
        _, out2, rep2 = self.run_scrub(make_pptx, "deck2.pptx", "--strip-notes")
        self.assert_gone(blob(out2), "SPEAKER-NOTE-ONE", "HIDDEN-NOTE-TWO")
        self.assertEqual(rep2["changes"]["notes_removed"], 2)

    def test_xlsx_text_cells_are_neutralised_and_formulas_kept(self):
        src, out, rep = self.run_scrub(make_xlsx, "book.xlsx")
        text = blob(out)
        self.assert_gone(text, "SECRET-COMMENT-XLSX", "THREADED-SECRET", "VML-COMMENT-SHAPE", "absPath", "D:\\Shared", "legacyDrawing", *TOKENS)
        for kept in ("<f>SUM(C1:C1)</f>", "<v>=text</v>", "Hello", "MSIP_Label_0001_Name", "x15ac"):
            self.assertIn(kept, text)
        for want in ("'=cmd|' /C calc'!A0", "'@rich", "'+1+1", "'-2+3", "'=bare"):
            self.assertIn(want, text)
        names = " ".join(zipfile.ZipFile(out).namelist())
        self.assert_gone(names, "comments1", "vmlDrawing", "threadedComment", "person")
        self.assertEqual(rep["changes"]["cells_neutralised"], 5)
        if openpyxl:
            ws = openpyxl.load_workbook(out)["Data"]
            self.assertEqual((ws["A1"].value, ws["A1"].data_type, ws["A2"].value), ("'=cmd|' /C calc'!A0", "s", "Hello"))
            self.assertEqual((ws["C1"].value, ws["C2"].value), (5, "=SUM(C1:C1)"))

    def test_second_scrub_changes_nothing(self):
        for maker, name in ((make_docx, "a.docx"), (make_pptx, "a.pptx"), (make_xlsx, "a.xlsx")):
            _, out, _ = self.run_scrub(maker, name)
            r = scrub.scrub(str(out), str(self.d / ("again-" + name)))
            self.assertFalse(r["changed"], name)

    def test_refusals_leave_no_output(self):
        ok = self.d / "ok.docx"
        make_docx(ok)
        cases = {"macro.docm": b"PK", "legacy.doc": b"\xd0\xcf\x11\xe0", "locked.docx": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "text.xlsx": b"just text", "note.txt": b"x"}
        for name, data in cases.items():
            (self.d / name).write_bytes(data)
            out = self.d / ("out-" + name)
            self.assertEqual(scrub.main([str(self.d / name), "--out", str(out)]), 1, name)
            self.assertFalse(out.exists() or Path(str(out) + ".tmp").exists(), name)
        self.assertEqual(scrub.main([str(ok), "--out", str(ok)]), 1)
        self.assertEqual(scrub.main([str(ok), "--out", str(self.d / "missing-dir" / "x.docx")]), 1)
        self.assertEqual(scrub.main([str(ok), "--out", str(self.d / "x.pptx")]), 1)

    @unittest.skipUnless(docx and openpyxl, "needs python-docx and openpyxl")
    def test_files_made_by_the_real_libraries(self):
        from openpyxl.comments import Comment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"] = "hello", "=1+1", "-5"
        ws["A3"].data_type = "s"
        ws["A1"].comment = Comment("NOTE-SECRET", "Jordan Example")
        wb.properties.creator = "Jordan Example"
        wb.save(self.d / "real.xlsx")
        doc = docx.Document()
        doc.add_comment(doc.add_paragraph("body text").runs[0], text="DOC-NOTE-SECRET", author="Jordan Example")
        doc.core_properties.author = "Jordan Example"
        doc.save(self.d / "real.docx")
        for name, secret in (("real.xlsx", "NOTE-SECRET"), ("real.docx", "DOC-NOTE-SECRET")):
            out = self.d / ("clean-" + name)
            self.assertTrue(scrub.scrub(str(self.d / name), str(out))["changed"])
            self.assert_gone(blob(out), secret, "Jordan Example")
            self.assert_valid(out)
        self.assertEqual([c.value for c in openpyxl.load_workbook(self.d / "clean-real.xlsx").active["A"]], ["hello", "=1+1", "'-5"])
        self.assertEqual(docx.Document(str(self.d / "clean-real.docx")).paragraphs[0].text, "body text")

    def test_lint_flags_missing_alt_text_titles_and_heading_jumps(self):
        deck = self.d / "lint.pptx"
        make_pptx(deck)
        self.assertEqual(sum("no title" in m for m in render.lint(str(deck))), 3)     # the fixture slides have no title placeholder
        pic = '<w:drawing><wp:inline xmlns:wp="x"><wp:docPr id="{n}" name="Picture"{alt}/></wp:inline></w:drawing>'
        body = ('<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr></w:p><w:p><w:pPr><w:pStyle w:val="Heading3"/></w:pPr></w:p>'
                "<w:p><w:r>" + pic.format(n=1, alt="") + pic.format(n=2, alt=' descr="A chart of bookings"') + "</w:r></w:p>")
        memo = self.d / "lint.docx"
        make(memo, "word/document.xml", {"word/document.xml": f'{XML}<w:document {W}><w:body>{body}</w:body></w:document>'},
             {"word/document.xml": CTYPE + "wordprocessingml.document.main+xml"})
        self.assertEqual(render.lint(str(memo)), ["picture without alt text", "heading level jumps from 1 to 3"])
        r = subprocess.run([sys.executable, str(SCRIPTS / "render.py"), str(memo), "--lint"], capture_output=True, text=True)
        self.assertEqual((r.returncode, r.stdout.count("LINT:")), (0, 2))


if __name__ == "__main__":
    unittest.main()

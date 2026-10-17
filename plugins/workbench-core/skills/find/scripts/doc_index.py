#!/usr/bin/env python3
"""doc_index.py: one local index over work libraries, captured content and the how-to library (the `find` skill).

  doc_index.py [--db P | --folder F] <command>

  discover [--json]                  libraries on this machine (read-only, never downloads a placeholder)
  add-root PATH [--label L]          register a folder    remove-root L [--purge]    roots [--json]
  index FOLDER [--rebuild] [--force] [--include-labelled]     add a root and index it
  refresh [--all | --root L] [--rebuild] [--embed auto|off] [--max-minutes N] [--detach]
  embed [--model M] [--batch N] [--max-minutes N]             compute vectors (downloads the model once)
  model import PATH | model export OUT.zip                    put the pinned model in place from a folder or .zip with no
                                                              network (export builds that .zip from the cached model)
  search "Q" [--mode auto|hybrid|keyword] [--limit N] [--root L] [--path S] [--source S] [--rerank]
  show "REF" [--context N]           files "WORDS"           related "REF|PATH" [--depth N]     newest "REF|PATH|WORDS"
  convert FILE --json                label FILE --json       status | report [--json]           --purge [--root L]

Three layers share one SQLite file: a catalog row for every file seen (name, path, title, mtime, size, label,
cloud-only), paragraph rows searched by BM25 and, when vectors exist, by a numpy dot product, fused with reciprocal
rank fusion; and explicit edges (folders, links, ticket keys, people, versions). No LLM, no network at query time.
The catalog commands (files, discover, roots, status, label, show, related, newest) use only the standard library.
Exit: 0 ok, 1 error, 2 usage.
"""
import argparse
import contextlib
import datetime
import fnmatch
import hashlib
import importlib.util
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import time
import zipfile
import zlib
import xml.etree.ElementTree as ET
from urllib.parse import urlparse

SCHEMA_VERSION = 3                      # 0.2.0 databases had user_version 0
MODEL = "BAAI/bge-small-en-v1.5"
MODEL_GLOB = "models--*bge-small-en-v1.5*"
MODEL_FILE = "model_optimized.onnx"
MODEL_SHA256 = "51f1bd0addd6e859e42c2c8021a5e5461385bb676a649f4b269aa445449f2431"   # Qdrant/bge-small-en-v1.5-onnx-Q model_optimized.onnx, as published
RRF_K, TOPK, FANOUT_CAP, STALE_HOURS, MODEL_RETRY_HOURS = 60, 50, 50, 24, 24
EXTS = {".docx", ".pptx", ".xlsx", ".pdf", ".html", ".htm", ".md", ".markdown", ".txt"}
PLAIN = {".md", ".markdown", ".txt"}
OFFICE = {".docx", ".pptx", ".xlsx"}
HIDDEN_KINDS = ("comments", "tracked_changes", "speaker_notes", "hidden_sheets", "hidden_slides")
# Windows file attributes that mark a cloud placeholder whose content is not on disk.
CLOUD_ONLY = 0x1000 | 0x40000 | 0x400000
MOUNT_POINT = 0xA0000003          # IO_REPARSE_TAG_MOUNT_POINT: an NTFS junction
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # an encrypted or password-protected Office file is an OLE container
MAX_PARA = 1800
SKIPDIRS = {"node_modules", "__pycache__", "$RECYCLE.BIN"}
STOP = frozenset("a an and are as at be by can do does for from how i if in is it its me my of on or our should that the "
                 "their there this to was we what when where which who why will with you your".split())

SCHEMA = """
create table meta(k text primary key, v text);
create table roots(id integer primary key, label text unique not null, path text not null, kind text not null);
create table files(id integer primary key, root_id int not null, rel text not null, ref text not null, dir text, name text,
  ext text, title text, headings text, hay text, size int, mtime real, sha256 text, status text not null, reason text,
  error text, label text, cloud_only int default 0, paragraphs int default 0, indexed_at text, url text, source text,
  fm_status text, modified text, captured int default 0, hidden text);
create unique index files_ref on files(ref);
create unique index files_rel on files(root_id, rel);
create index files_dir on files(root_id, dir);
create table paras(id integer primary key, file_id int not null, n int not null, heading text, locator text,
  hidden text, body text, thash text);
create index paras_file on paras(file_id);
create index paras_thash on paras(thash);
create virtual table fts using fts5(text, tokenize='porter unicode61 remove_diacritics 2');
create table vec(para_rowid int, model text, dim int, blob blob, primary key(para_rowid, model)) without rowid;
create table edges(src text, dst text, kind text, primary key(src, dst, kind)) without rowid;
create index edges_dst on edges(dst);
create table keys(path text, key text, primary key(path, key)) without rowid;
create index keys_key on keys(key);
create table linkraw(file_id int, kind text, target text);
create index linkraw_file on linkraw(file_id);
"""


# ───────────────────────── paths and database ─────────────────────────
def lad():
    return (os.environ.get("LOCALAPPDATA") if os.name == "nt" else None) or os.path.join(os.path.expanduser("~"), ".cache")


def wb_config():
    home = os.environ.get("CLAUDE_CONFIG_DIR") or os.path.join(os.path.expanduser("~"), ".claude")
    try:
        with open(os.path.join(home, "workbench", "config.json"), encoding="utf-8-sig") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def shared_db():
    return os.path.join(lad(), "doc-index", "knowledge", "index.sqlite")


def models_dir():
    return os.path.join(lad(), "doc-index", "models")


def resolve_db(db=None, folder=None):
    if db:
        return os.path.abspath(db)
    if folder:   # a per-folder index; never the 0.2.0 file name, so an old index is left alone
        key = hashlib.sha1(os.path.abspath(folder).lower().encode()).hexdigest()[:12]
        return os.path.join(lad(), "doc-index", key, "knowledge.sqlite")
    return os.path.abspath(os.environ.get("WORKBENCH_KNOWLEDGE_DB") or wb_config().get("knowledge_db") or shared_db())


def lp(p):
    """Absolute path; long Windows paths get the \\\\?\\ prefix so a deep captured tree still opens."""
    p = os.path.abspath(p)
    if os.name == "nt" and len(p) >= 240 and not p.startswith("\\\\?\\"):
        return "\\\\?\\UNC\\" + p[2:] if p.startswith("\\\\") else "\\\\?\\" + p
    return p


def connect(db):
    os.makedirs(os.path.dirname(os.path.abspath(db)), exist_ok=True)
    con = sqlite3.connect(db, timeout=30)
    con.execute("pragma busy_timeout=30000")
    con.execute("pragma journal_mode=wal")
    con.execute("pragma secure_delete=on")      # deleted plaintext is zeroed, not left in free pages
    ensure_schema(con, db)
    return con


def slug(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "-", s).strip("-.") or "root"


def unique_label(con, want):
    label, n = want, 1
    while con.execute("select 1 from roots where label=?", (label,)).fetchone():
        n += 1
        label = f"{want}-{n}"
    return label


def kind_of(label, path):
    parts = [p.lower() for p in re.split(r"[\\/]+", path)]
    if label in ("captured", "howto", "kb"):   # kb: the knowledge-base folder (notes the save skill writes)
        return label
    if any(p.startswith("onedrive -") for p in parts):
        return "onedrive-business"
    return "sharepoint-sync" if any("sharepoint" in p for p in parts) else "documents"


def long_path(p):
    """Windows: p with 8.3 short names (PROGRA~1) expanded to the long ones (GetLongPathNameW); a part that does not exist
    is kept as written. Elsewhere, and whenever the call fails, p itself."""
    if os.name != "nt" or "~" not in p:   # a short name always carries "~"; without one there is nothing to expand, and asking the
        return p                          # file system about an unreachable network share (\\host\share) stalls for ~20 s per call
    try:
        import ctypes
        buf = ctypes.create_unicode_buffer(32768)
        head, tail = p, []
        while True:
            if ctypes.windll.kernel32.GetLongPathNameW(head, buf, len(buf)):
                return os.path.join(buf.value, *reversed(tail))
            parent, name = os.path.split(head)
            if not name or parent == head:
                return p
            head, tail = parent, tail + [name]
    except Exception:   # noqa: BLE001
        return p


def norm_path(p):
    """A folder path as the OS compares it: absolute, short names expanded, no trailing separator, case-folded on Windows."""
    return os.path.normcase(long_path(os.path.abspath(p)))


def add_root(con, path, label=None):
    path = os.path.abspath(path)
    for lab, p in con.execute("select label, path from roots"):   # the same folder spelled another way is the same root
        if norm_path(p) == norm_path(path):
            return lab
    label = slug(label or os.path.basename(path.rstrip("\\/")))
    if con.execute("select 1 from roots where label=?", (label,)).fetchone():
        raise SystemExit(f"root label '{label}' is already used for another folder")
    con.execute("insert into roots(label, path, kind) values (?,?,?)", (label, path, kind_of(label, path)))
    con.commit()
    return label


def set_meta(con, k, v):
    con.execute("insert or replace into meta values (?,?)", (k, v))


def get_meta(con, k, default=None):
    r = con.execute("select v from meta where k=?", (k,)).fetchone()
    return r[0] if r else default


def add_note(con, text):
    notes = json.loads(get_meta(con, "notes", "[]"))
    if text not in notes:
        set_meta(con, "notes", json.dumps((notes + [text])[-5:]))
        con.commit()


def ensure_schema(con, db):
    if con.execute("pragma user_version").fetchone()[0] == SCHEMA_VERSION:
        return
    con.execute("begin immediate")
    try:
        if con.execute("pragma user_version").fetchone()[0] != SCHEMA_VERSION:   # another process may have won the race
            fresh = not con.execute("select 1 from sqlite_master where type='table' limit 1").fetchone()
            keep = []
            if not fresh:   # the index is a cache and is rebuilt; the registered roots are kept
                try:
                    keep = con.execute("select label, path, kind from roots").fetchall()
                except sqlite3.Error:
                    keep = []
                for (name,) in con.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%' "
                                           "order by sql like 'CREATE VIRTUAL%' desc").fetchall():
                    try:   # a virtual table goes first and takes its shadow tables with it
                        con.execute(f'drop table if exists "{name}"')
                    except sqlite3.OperationalError:
                        pass
            for stmt in SCHEMA.split(";"):
                if stmt.strip():
                    con.execute(stmt)
            con.executemany("insert into roots(label, path, kind) values (?,?,?)", keep)
            if fresh and os.path.abspath(db) == os.path.abspath(shared_db()):
                migrate_legacy(con)
            con.execute(f"pragma user_version={SCHEMA_VERSION}")
        con.execute("commit")
    except BaseException:
        con.execute("rollback")
        raise


def migrate_legacy(con):
    """First open of the shared database: every 0.2.0 index at <LAD>/doc-index/<12 hex>/index.sqlite whose indexed folder
    still exists becomes a registered root (re-indexed by the next refresh). The old files are left where they are."""
    base, moved = os.path.join(lad(), "doc-index"), []
    for d in sorted(os.listdir(base)) if os.path.isdir(base) else []:
        old = os.path.join(base, d, "index.sqlite")
        if not (re.fullmatch(r"[0-9a-f]{12}", d) and os.path.isfile(old)):
            continue
        try:
            o = sqlite3.connect(f"file:{old}?mode=ro", uri=True)
            try:
                root = dict(o.execute("select k, v from meta")).get("root")
            finally:
                o.close()
        except sqlite3.Error:
            continue
        if root and os.path.isdir(root) and not con.execute("select 1 from roots where path=?", (os.path.abspath(root),)).fetchone():
            label = unique_label(con, slug(os.path.basename(root.rstrip("\\/"))))
            con.execute("insert into roots(label, path, kind) values (?,?,?)", (label, os.path.abspath(root), kind_of(label, root)))
            moved.append(old)
    if moved:
        set_meta(con, "migrated_from", json.dumps(moved))


def iso(ts):
    return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ") if ts else None


def now_iso():
    return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# ───────────────────────── files on disk ─────────────────────────
def is_cloud_only(st):
    return bool(getattr(st, "st_file_attributes", 0) & CLOUD_ONLY)


def nested_roots(con, rid, path):
    """normcase'd absolute paths of the OTHER roots that sit strictly below this root. A nested root owns its folder: the
    outer root's walk skips it, so every file is indexed once, under the most specific root. A root with the same path is
    not nested, and neither is a sibling that only shares a name prefix (outer2 next to outer)."""
    base = norm_path(path)
    pre = base if base.endswith(os.sep) else base + os.sep
    return {p for p in (norm_path(r[0]) for r in con.execute("select path from roots where id<>?", (rid,))) if p != base and p.startswith(pre)}


def is_junction(e):
    """True for an NTFS junction (a mount-point reparse point) among os.scandir entries. A OneDrive placeholder is a
    reparse point too, under another tag, and is NOT one. Python 3.12+ answers this itself."""
    if hasattr(e, "is_junction"):
        return e.is_junction()
    st = e.stat(follow_symlinks=False)
    return bool(getattr(st, "st_file_attributes", 0) & 0x400) and getattr(st, "st_reparse_tag", MOUNT_POINT) == MOUNT_POINT


def walk(root, skip_underscore=False, skip_top=(), skip_paths=(), skipped=None):
    """(relpath, stat) for every file. DirEntry.stat() reads cached attributes on Windows, so it never opens (and never
    hydrates) a placeholder. skip_underscore (captured roots): a folder whose name starts with _ is a staging area.
    skip_top: folder names skipped directly under the root only (the kb root's daily/ save log).
    skip_paths: normcase'd absolute folders that belong to another root; their relpaths are appended to `skipped`.
    A junction is never entered: its files belong to the folder it points at (indexed there, or not at all)."""
    stack = [(os.path.abspath(root), "")]
    while stack:
        d, rel_dir = stack.pop()
        try:
            entries = list(os.scandir(lp(d)))
        except OSError:
            continue
        for e in entries:
            if e.name.startswith((".", "~$")) or e.name in SKIPDIRS or (not rel_dir and e.name in skip_top):
                continue
            rel = f"{rel_dir}/{e.name}" if rel_dir else e.name
            try:
                if e.is_dir(follow_symlinks=False):
                    sub = os.path.join(d, e.name)
                    if is_junction(e):
                        continue
                    if os.path.normcase(sub) in skip_paths:
                        if skipped is not None:
                            skipped.append(rel)
                    elif not (skip_underscore and e.name.startswith("_")):
                        stack.append((sub, rel))
                elif e.is_file(follow_symlinks=False):
                    yield rel, e.stat(follow_symlinks=False)
            except OSError:
                continue


def sha256(path):
    h = hashlib.sha256()
    with open(lp(path), "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


FM_LINE = re.compile(r"^([A-Za-z_][\w-]*):[ \t]*(.*)$")


def parse_front_matter(text):
    """(dict|None, body). Lines are `key: <JSON value>`; a value that is not JSON is kept as plain text. Unknown keys
    are carried along and ignored by the readers, never an error."""
    text = text.lstrip("\ufeff")
    lines = text.split("\n")
    if lines[0].strip() != "---":
        return None, text
    end = next((i for i in range(1, min(len(lines), 300)) if lines[i].strip() == "---"), None)
    if end is None:
        return None, text
    fm = {}
    for ln in lines[1:end]:
        m = FM_LINE.match(ln.rstrip("\r"))
        if m:
            v = m.group(2).strip()
            try:
                fm[m.group(1)] = json.loads(v)
            except ValueError:
                fm[m.group(1)] = v.strip("'\"")
    return fm, "\n".join(lines[end + 1:])


def read_text(path, cap=None):
    with open(lp(path), "rb") as f:
        data = f.read(cap) if cap else f.read()
    return data.decode("utf-8", "replace")


def is_captured(fm):
    return bool(fm) and "source" in fm and "captured_at" in fm


def fm_label(fm):
    labels = fm.get("labels") if is_captured(fm) else None
    labels = labels if isinstance(labels, list) else ([labels] if labels else [])
    return ", ".join(str(x) for x in labels) or None


# ───────────────────────── labels ─────────────────────────
MIP_NS = "{http://schemas.microsoft.com/office/2020/mipLabelMetadata}"   # [MS-OFFCRYPTO] section 2.6.4, Namespaces
LABEL_PART_CAP = 1024 * 1024   # label parts are a few hundred bytes; anything bigger is suspect


def _xsd_true(v):
    return str(v).strip().lower() in ("1", "true")


def safe_xml(data):
    """Parse XML from an Office part. Office parts never carry a DTD, so one is refused (no entity tricks)."""
    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
        raise ET.ParseError("DTD not allowed")
    return ET.fromstring(data)


def _read_part(z, info):
    """A part's bytes, or None when it is larger than LABEL_PART_CAP. The size is checked before any decompression, and
    the read itself is bounded too, so a zip bomb cannot exhaust memory."""
    if info.file_size > LABEL_PART_CAP:
        return None
    with z.open(info) as f:
        data = f.read(LABEL_PART_CAP + 1)
    return None if len(data) > LABEL_PART_CAP else data


def sensitivity_label(path):
    """The Office file's sensitivity label, or None. Reads exactly two parts, matched case-insensitively:
    docProps/custom.xml (MSIP_Label_<id>_Name) and docMetadata/LabelInfo.xml (co-authoring; any label element enabled
    and not removed counts). An oversized or unreadable part fails closed (counts as labelled)."""
    try:
        z = zipfile.ZipFile(lp(path))
    except (OSError, zipfile.BadZipFile):
        return None
    with z:
        parts = {i.filename.lower(): i for i in z.infolist()}
        name = None
        info = parts.get("docprops/custom.xml")
        if info is not None:
            data = _read_part(z, info)
            if data is None:
                return "label (oversized custom properties part)"
            m = re.search(r'name="MSIP_Label_[0-9a-fA-F-]+_Name"[^>]*>\s*<vt:lpwstr>([^<]*)<', data.decode("utf-8", "replace"))
            name = m.group(1) if m else None
        info = parts.get("docmetadata/labelinfo.xml")
        if info is not None:
            data = _read_part(z, info)
            if data is None:
                return name or "label (oversized LabelInfo part)"
            try:
                root = safe_xml(data)
            except ET.ParseError:
                return name or "label (unreadable LabelInfo part)"
            for el in root.iter(MIP_NS + "label"):
                if _xsd_true(el.get("enabled")) and not _xsd_true(el.get("removed")):
                    return name or f"label {el.get('id') or '?'}"
        return name


def pdf_label(path):
    """A PDF's sensitivity label (metadata keys starting MSIP_Label_), or None. Looks in the raw bytes and inside
    Flate streams (a modern PDF keeps its Info and XMP in compressed objects); stdlib only."""
    import zlib
    try:
        with open(lp(path), "rb") as f:
            data = f.read(64 << 20)
    except OSError:
        return None
    chunks, budget = [data], 16 << 20
    for m in re.finditer(rb"stream\r?\n(.*?)\r?\nendstream", data, re.S):
        try:
            d = zlib.decompressobj().decompress(m.group(1), 4 << 20)
        except zlib.error:
            continue
        budget -= len(d)
        chunks.append(d)
        if budget < 0:
            break
    guid = name = None
    for b in chunks:
        g = re.search(rb"MSIP_Label_([0-9a-fA-F-]{8,40})_\w+", b)
        if not g:
            continue
        guid = guid or g.group(1).decode()
        n = re.search(rb"MSIP_Label_[0-9a-fA-F-]+_Name\s*(?:\(([^)]*)\)|<([0-9A-Fa-f]+)>|>([^<]*)<)", b)
        if n and not name:
            raw = n.group(1) if n.group(1) is not None else (bytes.fromhex(n.group(2).decode()) if n.group(2) else n.group(3))
            name = (raw.decode("utf-16", "replace") if raw[:2] in (b"\xfe\xff", b"\xff\xfe") else raw.decode("latin-1")).strip() or None
    return (name or f"label {guid}") if guid else None


def file_label(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in OFFICE:
        return sensitivity_label(path)
    if ext == ".pdf":
        return pdf_label(path)
    if ext in PLAIN:
        try:
            return fm_label(parse_front_matter(read_text(path, 1 << 16))[0] or {})
        except OSError:
            return None
    return None


MAX_ZIP_MEMBERS = 10_000


def unzipped_too_big(path, cap_mb):
    """Reason string if an Office zip would expand past cap_mb or lists too many parts, else None. Uses the sizes in the
    zip directory only; nothing is decompressed."""
    try:
        with zipfile.ZipFile(lp(path)) as z:
            infos = z.infolist()
    except (OSError, zipfile.BadZipFile):
        return None   # not a readable zip: conversion reports the real error
    total = sum(i.file_size for i in infos)
    if total > cap_mb * 1024 * 1024 or len(infos) > MAX_ZIP_MEMBERS:
        return f"too large when uncompressed ({total // (1024 * 1024)} MB in {len(infos)} parts; cap {cap_mb} MB, {MAX_ZIP_MEMBERS} parts)"
    return None


# ───────────────────────── conversion ─────────────────────────
# A converted file is a list of segments (text, locator, hidden_kind). The hidden kind comes from the file's structure,
# never from marker text inside the document, so hidden content cannot un-flag itself.
_md = None


def _need(mod):
    if importlib.util.find_spec(mod) is None:
        raise RuntimeError(f"{mod} is not installed: pip install -r requirements.txt")


def _local(el):
    return el.tag.rsplit("}", 1)[-1]


def _texts(el):
    return " ".join(t.text.strip() for t in el.iter() if _local(t) in ("t", "text", "delText") and t.text and t.text.strip())


COMMENT_PART = re.compile(r"^(word/comments[^/]*\.xml|xl/comments[^/]*\.xml|xl/comments/[^/]+\.xml|xl/threadedComments/[^/]+\.xml|"
                          r"ppt/comments/[^/]+\.xml)$", re.I)


def office_hidden(path, docx):
    """([comment texts], tracked-change count, [deleted texts]) read from the zip with the stdlib, size-bounded."""
    comments, tracked, deleted = [], 0, []
    with zipfile.ZipFile(lp(path)) as z:
        for info in z.infolist():
            is_doc = docx and info.filename == "word/document.xml"
            if (COMMENT_PART.match(info.filename) or is_doc) and info.file_size <= (64 << 20):
                try:
                    root = safe_xml(z.read(info))
                except ET.ParseError:
                    continue
                if is_doc:
                    for el in root.iter():
                        if _local(el) in ("ins", "del", "moveFrom", "moveTo"):
                            tracked += 1
                            if _local(el) in ("del", "moveFrom") and _texts(el):
                                deleted.append(_texts(el))
                else:
                    for el in root.iter():
                        if _local(el) in ("comment", "cm", "threadedComment"):
                            comments.append(_texts(el))
    return comments, tracked, deleted


class _Visible:
    """HTML to text with the invisible parts dropped: comments, script, style, and any element hidden by the hidden
    attribute or an inline display/visibility/opacity/font-size rule."""
    def __init__(self):
        from html.parser import HTMLParser
        outer = self

        class P(HTMLParser):
            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.out, self.stack, self.skip = [], [], None

            def handle_starttag(self, tag, attrs):
                a = dict(attrs)
                if tag in outer.VOID:
                    if self.skip is None and tag == "br":
                        self.out.append("\n")
                    return
                self.stack.append(tag)
                style = (a.get("style") or "").replace(" ", "").lower()
                if self.skip is None and (tag in ("script", "style", "template") or "hidden" in a or
                                          re.search(r"display:none|visibility:hidden|opacity:0(?![.\d])|font-size:0(?![.\d])", style)):
                    self.skip = len(self.stack)
                if self.skip is None:
                    if tag in outer.HEAD:
                        self.out.append("\n\n" + "#" * int(tag[1]) + " ")
                    elif tag in outer.BLOCK:
                        self.out.append("\n\n")

            def handle_endtag(self, tag):
                if tag in self.stack:
                    del self.stack[len(self.stack) - self.stack[::-1].index(tag) - 1:]
                if self.skip is not None and len(self.stack) < self.skip:
                    self.skip = None
                if self.skip is None and tag in outer.BLOCK | outer.HEAD:
                    self.out.append("\n\n")

            def handle_data(self, data):
                if self.skip is None:
                    self.out.append(data)

        self.parser = P()

    VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "track", "wbr"}
    HEAD = {"h1", "h2", "h3", "h4", "h5", "h6"}
    BLOCK = {"p", "div", "li", "tr", "section", "article", "header", "footer", "ul", "ol", "table", "blockquote", "pre"}

    def text(self, html):
        self.parser.feed(html)
        self.parser.close()
        t = re.sub(r"[ \t\r\f\v]+", " ", "".join(self.parser.out))
        return re.sub(r"\n[ ]*", "\n", re.sub(r"\n{3,}", "\n\n", t)).strip()


def _pptx_shapes(shapes):
    for sh in shapes:
        if hasattr(sh, "shapes"):
            yield from _pptx_shapes(sh.shapes)
        else:
            yield sh


def convert_segments(path):
    """(segments, hidden counts, title). Raises on an unreadable file."""
    ext = os.path.splitext(path)[1].lower()
    hidden = dict.fromkeys(HIDDEN_KINDS, 0)
    title = None
    if ext in PLAIN:
        fm, body = parse_front_matter(read_text(path))
        return [(body, None, None)], hidden, (str(fm["title"]) if fm and fm.get("title") else None)
    with open(lp(path), "rb") as f:
        head = f.read(8)
    if ext in OFFICE and head == OLE_MAGIC:
        raise ValueError("encrypted, password-protected or rights-managed (OLE container, not an Office zip)")
    if ext in (".html", ".htm"):
        return [(_Visible().text(read_text(path)), None, None)], hidden, None
    segs = []
    if ext == ".pdf":
        _need("pdfminer")
        from pdfminer.high_level import extract_text
        for i, page in enumerate(extract_text(lp(path)).split("\f"), 1):
            if page.strip():
                segs.append((page, f"page {i}", None))
        first = segs[0][0].strip().split("\n")[0].strip() if segs else ""
        return segs, hidden, (first if 0 < len(first) <= 100 else None)   # the metadata Title is never read
    if ext == ".docx":
        global _md
        if _md is None:
            _need("markitdown")
            from markitdown import MarkItDown
            _md = MarkItDown()
        segs.append((_md.convert(lp(path)).text_content or "", None, None))
        comments, hidden["tracked_changes"], deleted = office_hidden(path, True)
        hidden["comments"] = len(comments)
        segs += [(t, None, "comments") for t in comments if t] + [(t, None, "tracked_changes") for t in deleted if t]
    elif ext == ".pptx":
        _need("pptx")
        from pptx import Presentation
        prs = Presentation(lp(path))
        for i, slide in enumerate(prs.slides, 1):
            hid_slide = slide._element.get("show") == "0"
            tshape = slide.shapes.title
            body = []
            for sh in _pptx_shapes(slide.shapes):
                if sh.has_text_frame and (tshape is None or sh.shape_id != tshape.shape_id) and sh.text_frame.text.strip():
                    body.append(sh.text_frame.text.strip())
                elif getattr(sh, "has_table", False) and sh.has_table:
                    body.append("\n".join("| " + " | ".join(c.text.strip() for c in r.cells) + " |" for r in sh.table.rows))
            stitle = tshape.text_frame.text.strip() if tshape is not None and tshape.has_text_frame else ""
            title = title or stitle or None
            segs.append((("## " + stitle + "\n\n" if stitle else "") + "\n\n".join(body), f"slide {i}", "hidden_slides" if hid_slide else None))
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ""
            if notes:
                hidden["speaker_notes"] += 1
                segs.append((notes, f"slide {i}", "hidden_slides" if hid_slide else "speaker_notes"))
            hidden["hidden_slides"] += hid_slide
        comments = office_hidden(path, False)[0]
        hidden["comments"] = len(comments)
        segs += [(t, None, "comments") for t in comments if t]
    elif ext == ".xlsx":
        _need("openpyxl")
        import openpyxl
        wb = openpyxl.load_workbook(lp(path), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "/") for c in r]
                        for r in ws.iter_rows(max_row=20000, values_only=True)]   # ponytail: first 20,000 rows per sheet
                rows = [r for r in rows if any(c.strip() for c in r)]
                hid_sheet = ws.sheet_state != "visible"
                hidden["hidden_sheets"] += hid_sheet
                if not rows:
                    continue
                fmt = lambda r: "| " + " | ".join(r) + " |"   # noqa: E731
                head_row, blocks, buf = fmt(rows[0]), [], []
                for line in map(fmt, rows[1:]):   # blocks of rows, each led by the header row
                    if buf and sum(len(x) + 1 for x in buf) + len(line) > MAX_PARA:
                        blocks.append("\n".join([head_row] + buf))
                        buf = []
                    buf.append(line)
                blocks.append("\n".join([head_row] + buf))
                segs.append((f"## {ws.title}\n\n" + "\n\n".join(blocks), f"sheet {ws.title}", "hidden_sheets" if hid_sheet else None))
        finally:
            wb.close()
        comments = office_hidden(path, False)[0]
        hidden["comments"] = len(comments)
        segs += [(t, None, "comments") for t in comments if t]
    else:
        raise ValueError(f"unsupported type {ext}")
    return segs, hidden, title


def convert_doc(path, include_labelled=False):
    out = {"path": path, "label": None, "segments": [], "hidden_content": dict.fromkeys(HIDDEN_KINDS, 0), "title": None,
           "fm": None, "error": None}
    try:
        out["label"] = file_label(path)
        if out["label"] and not include_labelled:
            return out
        out["segments"], out["hidden_content"], out["title"] = convert_segments(path)
        if os.path.splitext(path)[1].lower() in PLAIN:
            out["fm"] = parse_front_matter(read_text(path))[0]
    except Exception as e:   # one bad file must never stop a run
        out["error"] = f"{type(e).__name__}: {str(e)[:300]}"
    return out


HEAD = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


def paragraphs(segments):
    """[(heading path, text, locator, hidden kind)]: blank-line blocks and Markdown headings, in order. A heading is its
    own paragraph and extends the heading path; blocks longer than MAX_PARA are split on line boundaries."""
    out, stack = [], []
    for text, loc, hid in segments:
        buf = []

        def flush():
            if buf:
                out.append((" > ".join(s[1] for s in stack), "\n".join(buf).strip(), loc, hid))
                buf.clear()
        for line in text.replace("\r\n", "\n").split("\n"):
            if not line.strip():
                flush()
                continue
            m = HEAD.match(line)
            if m:
                flush()
                stack[:] = [s for s in stack if s[0] < len(m.group(1))] + [(len(m.group(1)), m.group(2))]
                buf.append(line.strip())
                flush()
                continue
            while len(line) > 2 * MAX_PARA:   # one enormous line (minified text): cut at a space
                cut = line.rfind(" ", 0, MAX_PARA)
                cut = cut if cut > 0 else MAX_PARA
                flush() if buf else None
                buf.append(line[:cut].strip())
                flush()
                line = line[cut:].strip()
            if buf and sum(len(x) + 1 for x in buf) + len(line) > MAX_PARA:
                flush()
            buf.append(line)
        flush()
    return [p for p in out if p[1]]


# ───────────────────────── keys and links ─────────────────────────
TICKET = re.compile(r"\b([A-Z][A-Z0-9]{1,9}-\d{1,6})\b")
CODE = re.compile(r"\b([A-Z]{2,5}\d{3,6})\b")
ACRONYM = re.compile(r"\b(?:[A-Z][a-z]+\s+){1,6}\(([A-Z]{2,8})\)")
NOT_KEYS = {"UTF", "SHA", "ISO", "AES", "RSA", "RFC", "CVE", "TLS", "SSL", "MD", "IPV", "TCP", "UDP", "HTTP", "PDF", "ASCII"}
MDLINK = re.compile(r"\[(?:[^\]]*)\]\(\s*<?([^)\s>]+)>?(?:\s+\"[^\"]*\")?\s*\)")
BARE = re.compile(r"https?://[^\s)>\]\"']+")
WIKI = re.compile(r"\[\[([^\[\]\n]{1,300})\]\]")


def extract_keys(*texts):
    keys = set()
    for t in texts:
        for k in TICKET.findall(t) + CODE.findall(t):
            if k.split("-")[0].rstrip("0123456789") not in NOT_KEYS:
                keys.add("ticket:" + k)
        keys.update("acronym:" + a for a in ACRONYM.findall(t) if a not in NOT_KEYS)
    return sorted(keys)[:200]


def norm_person(s):
    s = str(s).strip()
    m = re.search(r"<([^<>]+@[^<>]+)>", s)
    return (m.group(1) if m else s).strip().strip("'\"").lower()


def fm_keys(fm):
    keys = set()

    def people(v):
        if isinstance(v, list):
            return [str(x) for x in v]
        if isinstance(v, str):
            return re.split(r"[;,]" if "@" in v else r";", v)
        return []
    for field in ("from", "to", "cc", "organizer", "attendees"):
        keys.update("person:" + norm_person(s) for s in people(fm.get(field)) if s.strip())
    for field, pre in (("thread_id", "thread:"), ("meeting_id", "meeting:")):
        if fm.get(field):
            keys.add(pre + str(fm[field]).strip().lower())
    return sorted(keys)


def extract_links(text):
    out = [("hyperlink", t) for t in MDLINK.findall(text) + BARE.findall(text)]
    return out + [("wiki", t) for t in WIKI.findall(text)]


def norm_url(u):
    u = re.split(r"[?#]", u.strip(), maxsplit=1)[0].rstrip("/")
    m = re.match(r"(https?://)([^/]+)(.*)$", u, re.I)
    return (m.group(1).lower() + m.group(2).lower() + m.group(3)) if m else u


VERSION = re.compile(r"^(?P<base>.+?)[ _.\-]+(?:v|ver|version|rev)[ _.\-]?(?P<v>\d+(?:\.\d+){0,2})(?:[ _.\-]+(?:final|draft|clean))?$", re.I)
DATED = re.compile(r"^(?P<base>.+?)[ _\-]+(?P<d>20\d{2}[-_.]?\d{2}[-_.]?\d{2})$")


def version_of(stem):
    """(base, kind, sort key) when the file name carries a version or a date, else None."""
    m = VERSION.match(stem)
    if m:
        return m.group("base").lower(), "v", tuple(int(x) for x in m.group("v").split("."))
    m = DATED.match(stem)
    return (m.group("base").lower(), "d", (int(re.sub(r"\D", "", m.group("d"))),)) if m else None


def resolve_edges(con, rid):
    """Recompute this root's explicit edges (hyperlink, wiki, supersedes) once all its files are catalogued."""
    rows = con.execute("select id, rel, ref, dir, name, url from files where root_id=?", (rid,)).fetchall()
    con.execute("delete from edges where kind in ('hyperlink','wiki','supersedes') and src in "
                "(select ref from files where root_id=?)", (rid,))
    by_rel = {r[1].lower(): r[2] for r in rows}
    by_stem = {}
    for _id, rel, ref, d, name, _u in rows:
        by_stem.setdefault(os.path.splitext(name)[0].lower(), []).append((d or "", ref))
    by_url = {norm_url(u): r for r, u in con.execute("select ref, url from files where url is not null")}
    ids = {r[0]: r for r in rows}
    import posixpath
    from urllib.parse import unquote
    new = set()
    for fid, kind, target in con.execute("select l.file_id, l.kind, l.target from linkraw l join files f on f.id=l.file_id "
                                         "where f.root_id=?", (rid,)).fetchall():
        _i, _rel, src, sdir, _n, _u = ids[fid]
        dst = None
        if kind == "hyperlink":
            if re.match(r"https?://", target, re.I):
                dst = by_url.get(norm_url(target))
            elif not re.match(r"[a-z][a-z0-9+.-]*:", target, re.I) and not target.startswith("#"):
                rel = posixpath.normpath(posixpath.join(sdir or "", unquote(re.split(r"[?#]", target, maxsplit=1)[0]))).lower()
                dst = by_rel.get(rel)
        else:
            t = target.split("|", 1)[0].split("#", 1)[0].strip().replace("\\", "/")
            dst = by_rel.get(t.lower()) or by_rel.get(t.lower() + ".md")
            if not dst:
                cands = by_stem.get(os.path.splitext(t.rsplit("/", 1)[-1])[0].lower(), [])
                dst = next((r for d, r in cands if d == (sdir or "")), cands[0][1] if cands else None)
        if dst and dst != src:
            new.add((src, dst, kind))
    groups, plain = {}, {}
    for _id, rel, ref, d, name, _u in rows:
        stem, ext = os.path.splitext(name)
        v = version_of(stem)
        if v:
            groups.setdefault((d, v[0], ext.lower(), v[1]), []).append((v[2], ref))
        else:
            plain[(d, stem.lower(), ext.lower())] = ref
    for (d, base, ext, _k), members in groups.items():
        if (d, base, ext) in plain:   # an unmarked sibling counts as the first version
            members.append(((-1,), plain[(d, base, ext)]))
        if len(members) > 1:
            members.sort()
            new.update((members[-1][1], ref, "supersedes") for _k2, ref in members[:-1])
    con.executemany("insert or ignore into edges values (?,?,?)", sorted(new))


# ───────────────────────── indexing ─────────────────────────
def ctx_text(title, hpath, name):
    """The contextual header (title > heading path > file name) that leads both the full-text and the embedding text."""
    return " > ".join(x for x in (title, hpath, name) if x)


def purge_files(con, where, args, incoming=True):
    """Delete every row that belongs to the selected files: paragraphs, full-text rows, vectors, keys, edges, links.
    incoming=False keeps the edges other files point at them with (a re-index, not a removal)."""
    ids = [r[0] for r in con.execute(f"select id from files where {where}", args)]
    refs = [r[0] for r in con.execute(f"select ref from files where {where}", args)]
    for chunk in (ids[i:i + 500] for i in range(0, len(ids), 500)):
        q = ",".join("?" * len(chunk))
        pids = f"select id from paras where file_id in ({q})"
        con.execute(f"delete from vec where para_rowid in ({pids})", chunk)
        con.execute(f"delete from fts where rowid in ({pids})", chunk)
        con.execute(f"delete from paras where file_id in ({q})", chunk)
        con.execute(f"delete from linkraw where file_id in ({q})", chunk)
    for chunk in (refs[i:i + 500] for i in range(0, len(refs), 500)):
        q = ",".join("?" * len(chunk))
        con.execute(f"delete from keys where path in ({q})", chunk)
        con.execute(f"delete from edges where src in ({q})", chunk)
        if incoming:
            con.execute(f"delete from edges where dst in ({q})", chunk)
    return ids


def forget_content(con, fid, ref):
    purge_files(con, "id=?", (fid,), incoming=False)


def store_doc(con, fid, ref, rel, name, doc):
    paras = paragraphs(doc["segments"])
    fm = doc["fm"] or {}
    heads = [re.sub(r"^#+\s*", "", p[1]) for p in paras if HEAD.match(p[1].split("\n")[0]) and p[1].count("\n") == 0]
    heads = [] if name.lower().endswith(".xlsx") else heads   # a sheet name is not a document title
    title = (str(fm["title"]) if fm.get("title") else None) or doc["title"] or (heads[0] if heads else None) or os.path.splitext(name)[0]
    forget_content(con, fid, ref)
    for n, (hp, body, loc, hid) in enumerate(paras, 1):
        text = ctx_text(title, hp, name) + "\n" + body
        pid = con.execute("insert into paras(file_id, n, heading, locator, hidden, body, thash) values (?,?,?,?,?,?,?)",
                          (fid, n, hp, loc, hid, body, hashlib.sha1(text.encode("utf-8", "replace")).hexdigest())).lastrowid
        con.execute("insert into fts(rowid, text) values (?,?)", (pid, text))
    visible = "\n".join(s[0] for s in doc["segments"] if not s[2])[:2_000_000]
    keys = set(extract_keys(name, title, visible)) | (set(fm_keys(fm)) if is_captured(fm) else set())
    con.executemany("insert or ignore into keys values (?,?)", [(ref, k) for k in sorted(keys)])
    con.executemany("insert into linkraw values (?,?,?)", [(fid, k, t[:500]) for k, t in extract_links(visible)[:500]])
    status = (fm.get("status") if fm else None)
    con.execute("update files set title=?, headings=?, hay=?, status='ok', reason=NULL, error=NULL, label=?, paragraphs=?, indexed_at=?,"
                " url=?, source=?, fm_status=?, modified=?, captured=?, hidden=? where id=?",
                (title, " | ".join(heads[:30])[:600], f"{name} {ref} {title} {' '.join(heads[:30])}".lower()[:2000], doc["label"],
                 len(paras), now_iso(), fm.get("url") if is_captured(fm) else None, fm.get("source") if is_captured(fm) else None,
                 str(status).lower() if status else None, fm.get("modified") if is_captured(fm) else None, int(is_captured(fm)),
                 json.dumps(doc["hidden_content"]), fid))
    return len(paras)


def mark(con, fid, status, reason, error, label=None):
    con.execute("update files set status=?, reason=?, error=?, label=?, paragraphs=0, indexed_at=? where id=?",
                (status, reason, error, label, now_iso(), fid))


def drop_files(con, rid, rels):
    for chunk in (rels[i:i + 500] for i in range(0, len(rels), 500)):
        where = "root_id=? and rel in (%s)" % ",".join("?" * len(chunk))
        purge_files(con, where, [rid, *chunk])
        con.execute(f"delete from files where {where}", [rid, *chunk])


def index_root(con, root, rebuild=False, force=False, include_labelled=False, max_mb=50, max_unzipped_mb=200,
               deadline=None, prune=False, log=print):
    rid, label, path = root
    path = long_path(path)   # walk from the long form, so a nested root's folder is recognised however it was spelled
    if not os.path.isdir(path):
        raise SystemExit(f"not a folder: {path}")
    stats = dict(new=0, updated=0, unchanged=0, failed=0, labelled=0, cloud_only=0, removed=0, catalogued=0, added=0, partial=False)
    kind = con.execute("select kind from roots where id=?", (rid,)).fetchone()[0]
    known = {r[0]: r for r in con.execute("select rel, id, size, mtime, status, cloud_only, label from files where root_id=?", (rid,))}
    seen, adds, changes, skipped = set(), [], [], []
    for rel, st in walk(path, skip_underscore=(kind == "captured"), skip_top=("daily",) if kind == "kb" else (),
                        skip_paths=nested_roots(con, rid, path), skipped=skipped):
        seen.add(rel)
        old = known.get(rel)
        cloud = int(is_cloud_only(st))
        stats["cloud_only"] += cloud
        ext = os.path.splitext(rel)[1].lower()
        status = ("cloud-only" if cloud else "pending") if ext in EXTS else "catalogued"
        if old is None:
            name = rel.rsplit("/", 1)[-1]
            adds.append((rid, rel, f"{label}/{rel}", rel.rpartition("/")[0], name, ext, os.path.splitext(name)[0],
                         f"{name} {label}/{rel} {os.path.splitext(name)[0]}".lower(), st.st_size, st.st_mtime, status, cloud))
        elif old[2] != st.st_size or old[3] != st.st_mtime or force or rebuild:
            changes.append((st.st_size, st.st_mtime, status, cloud, old[1]))
        else:   # same size and time: only the cloud flag or the label policy can change what happens next
            ns = old[4]
            if old[4] == "cloud-only" and not cloud:
                ns = status
            elif cloud and old[4] == "pending":
                ns = "cloud-only"
            elif (old[4] == "ok" and old[6] and not include_labelled) or (old[4] == "labelled" and include_labelled):
                ns = "pending"   # re-checked below: the label policy changed
            if ns != old[4] or old[5] != cloud:
                changes.append((st.st_size, st.st_mtime, ns, cloud, old[1]))
            else:
                stats["unchanged"] += 1
    for i in range(0, len(adds), 5000):
        con.executemany("insert into files(root_id, rel, ref, dir, name, ext, title, hay, size, mtime, status, cloud_only) "
                        "values (?,?,?,?,?,?,?,?,?,?,?,?)", adds[i:i + 5000])
    con.executemany("update files set size=?, mtime=?, status=?, cloud_only=? where id=?", changes)
    stats["added"], stats["catalogued"] = len(adds), len(seen)
    gone = [r for r in known if r not in seen]
    under = tuple(x.lower() + "/" for x in skipped)
    shadowed = [r for r in gone if r.lower().startswith(under)]   # indexed by an older version before a nested root took the folder over
    gone = [r for r in gone if not r.lower().startswith(under)]
    drop_files(con, rid, shadowed)   # not "gone": the nested root has them, so a duplicate is dropped even without prune
    stats["removed"] = len(shadowed)
    if rebuild or prune:
        if not seen and len(known) > len(shadowed):
            log("refusing to prune: no files found (is the folder offline or unmounted?)")
        elif gone:
            drop_files(con, rid, gone)
            stats["removed"] += len(gone)
    elif gone:
        log(f"{len(gone)} indexed file(s) no longer present; run with --rebuild to remove them")
    con.commit()
    # content pass
    where = "root_id=? and (status='pending'" + (" or status in ('failed','skipped')" if force else "") + ")"
    todo = con.execute(f"select id, rel, ref, name, size, sha256, paragraphs, status from files where {where} order by rel", (rid,)).fetchall()
    for fid, rel, ref, name, size, old_sha, old_paras, old_status in todo:
        if deadline and time.monotonic() > deadline:
            stats["partial"] = True
            break
        full, ext = os.path.join(path, rel), os.path.splitext(rel)[1].lower()
        lab = file_label(full)
        if lab and not include_labelled:   # never copy labelled content into the plaintext index by default
            forget_content(con, fid, ref)
            mark(con, fid, "labelled", "labelled", f"sensitivity label '{lab}': skipped (index it only with --include-labelled, where policy allows)", lab)
            stats["labelled"] += 1
            con.commit()
            continue
        bomb = unzipped_too_big(full, max_unzipped_mb) if ext in OFFICE else None
        if bomb or size > max_mb * 1024 * 1024:   # checked before any part is decompressed or handed to a converter
            forget_content(con, fid, ref)
            mark(con, fid, "skipped", "too-large-unzipped" if bomb else "too-large", bomb or f"larger than {max_mb} MB")
            stats["failed"] += 1
            con.commit()
            continue
        try:
            digest = sha256(full)
        except OSError as e:
            mark(con, fid, "failed", "convert-error", f"{type(e).__name__}: {str(e)[:300]}")
            stats["failed"] += 1
            con.commit()
            continue
        if old_paras and digest == old_sha and not force and old_status == "pending":   # touched, not changed
            con.execute("update files set status='ok', label=? where id=?", (lab, fid))
            stats["unchanged"] += 1
            con.commit()
            continue
        doc = convert_doc(full, include_labelled)
        con.execute("update files set sha256=? where id=?", (digest, fid))
        if doc["error"]:
            forget_content(con, fid, ref)
            enc = "encrypted" in doc["error"]
            mark(con, fid, "failed", "encrypted" if enc else "convert-error", doc["error"], lab)
            stats["failed"] += 1
        elif ext == ".pdf" and not doc["segments"]:
            forget_content(con, fid, ref)
            mark(con, fid, "skipped", "no-text-layer", "no text layer (scanned image?); OCR is out of scope", lab)
            stats["failed"] += 1
        else:
            doc["label"] = lab
            store_doc(con, fid, ref, rel, name, doc)
            stats["updated" if old_paras else "new"] += 1
        con.commit()
    if stats["added"] or stats["removed"] or force or rebuild or todo:
        resolve_edges(con, rid)
        con.commit()
    return stats


def index_stats_line(stats, db):
    return " · ".join(f"{k} {v}" for k, v in stats.items() if k != "partial") + f"  (index: {db})"


def coverage(con):
    skipped, content = {}, 0
    for status, reason, n in con.execute("select status, reason, count(*) from files group by status, reason"):
        if status == "ok":
            content += n
        else:
            key = {"catalogued": "unsupported-type", "pending": "not-yet-indexed"}.get(status) or reason or status
            skipped[key] = skipped.get(key, 0) + n
    return {"catalogued": con.execute("select count(*) from files").fetchone()[0], "content_indexed": content, "skipped": skipped}


# ───────────────────────── model and embeddings ─────────────────────────
class ModelUnavailable(Exception):
    pass


class ModelMissing(ModelUnavailable):
    """No model on disk and no download asked for."""


class ModelDownloadFailed(ModelUnavailable):
    """A download that was attempted and failed (as opposed to a missing package or HF_HUB_OFFLINE). Names the host and
    the error class, never the library's message: that wording talks about tokens although none is involved."""
    def __init__(self, host, error):
        super().__init__(f"model download from {host} failed ({error})")
        self.host, self.error = host, error


def hub_host():
    return urlparse(os.environ.get("HF_ENDPOINT") or "https://huggingface.co").hostname or "huggingface.co"


def error_name(e):
    """Error class, plus the HTTP status when the exception carries a response: 'HfHubHTTPError, HTTP 403'."""
    code = getattr(getattr(e, "response", None), "status_code", None)
    return type(e).__name__ + (f", HTTP {code}" if code else "")


@contextlib.contextmanager
def one_download_attempt(cap):
    """fastembed 0.9 retries a failed download three times, sleeping 3, 9 and 27 s between the tries, and logs every
    failure at ERROR in the hub's own words (they mention tokens). That is for a long-running service; here a failure is
    remembered for MODEL_RETRY_HOURS instead. So: one network attempt (the first error is kept in cap["exc"] and replayed
    to the library's later tries), no sleeping, no library log lines. Best effort: if fastembed's internals differ, the
    download just runs the way it always did."""
    try:
        import logging
        import warnings
        from fastembed.common import model_management as mm
        from loguru import logger
        mgmt, real_time = mm.ModelManagement, mm.time
        orig = mgmt.__dict__["download_files_from_huggingface"]
        func = orig.__func__ if isinstance(orig, classmethod) else None   # anything else: leave the library alone
    except Exception:   # noqa: BLE001
        func = None
    if func is None:
        yield
        return

    def once(cls, *a, **kw):
        if kw.get("local_files_only"):   # the library's cache probe makes no network request
            return func(cls, *a, **kw)
        if "exc" in cap:
            raise cap["exc"]
        try:
            return func(cls, *a, **kw)
        except Exception as e:   # noqa: BLE001
            cap["exc"] = e
            raise

    class NoSleep:
        sleep = staticmethod(lambda _s: None)

        def __getattr__(self, name):
            return getattr(real_time, name)
    hub_log = logging.getLogger("huggingface_hub")
    level = hub_log.level
    mgmt.download_files_from_huggingface = classmethod(once)
    mm.time = NoSleep()
    hub_log.setLevel(logging.CRITICAL)
    logger.disable("fastembed")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            yield
    finally:
        logger.enable("fastembed")
        hub_log.setLevel(level)
        mm.time = real_time
        mgmt.download_files_from_huggingface = orig


_EMB = {}


def find_model_file():
    import glob
    hits = glob.glob(os.path.join(models_dir(), MODEL_GLOB, "snapshots", "*", MODEL_FILE))
    return hits[0] if hits else None


def verified_model_file():
    """The model's ONNX path after its SHA-256 matched the pin (a stat-keyed marker skips re-hashing), else None. A file
    that does not match is deleted with its cache folder."""
    f = find_model_file()
    if not f:
        return None
    st = os.stat(f)
    marker = os.path.join(models_dir(), "verified.json")
    stamp = f"{st.st_size}:{st.st_mtime_ns}:{MODEL_SHA256}"
    try:
        with open(marker, encoding="utf-8") as m:
            if json.load(m).get(f) == stamp:
                return f
    except (OSError, ValueError):
        pass
    if sha256(f) != MODEL_SHA256:
        import shutil
        shutil.rmtree(os.path.dirname(os.path.dirname(os.path.dirname(f))), ignore_errors=True)
        raise ModelUnavailable("the model file failed its SHA-256 check and was deleted; keyword search is used until it is downloaded again")
    with open(marker, "w", encoding="utf-8") as m:
        json.dump({f: stamp}, m)
    return f


def offline():
    return os.environ.get("HF_HUB_OFFLINE", "").upper() in ("1", "TRUE", "YES", "ON")


def get_embedder(download=False):
    """A fastembed TextEmbedding for MODEL. Without `download` it never touches the network."""
    if "e" in _EMB:
        return _EMB["e"]
    f = verified_model_file()
    for mod in ("numpy", "fastembed"):
        if importlib.util.find_spec(mod) is None:
            raise ModelUnavailable(f"{mod} is not installed")
    if f is None:
        if not download:
            raise ModelMissing("the embedding model is not downloaded yet (run: doc_index.py embed)")
        if offline():
            raise ModelUnavailable("offline (HF_HUB_OFFLINE is set) and the embedding model is not downloaded")
    os.environ.setdefault("HF_HUB_DISABLE_TELEMETRY", "1")
    try:
        import truststore
        truststore.inject_into_ssl()
    except Exception:   # noqa: BLE001 - the model download then uses the default certificate store
        pass
    from fastembed import TextEmbedding
    os.makedirs(models_dir(), exist_ok=True)

    def make():
        return TextEmbedding(MODEL, cache_dir=models_dir(), lazy_load=True, local_files_only=f is not None or not download)
    fresh, cap = download and f is None, {}   # fresh: a real download, not the load of a cached model
    try:
        try:
            with one_download_attempt(cap) if fresh else contextlib.nullcontext():
                emb = make()
        except Exception as e:   # noqa: BLE001
            if not (fresh and re.search(r"ssl|certificate|tls", f"{e} {e.__cause__} {cap.get('exc')}", re.I)):
                raise
            os.environ["HF_HUB_DISABLE_XET"] = "1"   # the hf_xet path failed a TLS check: retry once over plain HTTPS
            import huggingface_hub.constants as hc
            hc.HF_HUB_DISABLE_XET = True
            cap.clear()
            with one_download_attempt(cap):
                emb = make()
        if verified_model_file() is None:
            raise ModelDownloadFailed(hub_host(), "files missing after the download") if fresh else ModelUnavailable(
                "the embedding model files are missing after the download")
    except ModelUnavailable:
        raise
    except Exception as e:   # noqa: BLE001
        if fresh:
            raise ModelDownloadFailed(hub_host(), error_name(cap.get("exc", e)))
        raise ModelUnavailable(f"embedding model unavailable ({type(e).__name__}: {str(e)[:160]})")
    _EMB["e"] = emb
    return emb


def vec_of(arr):
    import numpy as np
    a = np.asarray(arr, dtype=np.float32)
    return a / (np.linalg.norm(a, axis=-1, keepdims=True) + 1e-12)


def embed_pending(con, emb, batch=32, deadline=None):
    """Compute vectors for every paragraph that has none. Identical text (copies across library versions) is embedded once."""
    done = reused = 0
    while True:
        rows = con.execute("select p.id, p.thash, f.title, p.heading, f.name, p.body from paras p join files f on f.id=p.file_id "
                           "where not exists (select 1 from vec v where v.para_rowid=p.id and v.model=?) limit 256", (MODEL,)).fetchall()
        if not rows or (deadline and time.monotonic() > deadline):
            return done, reused, bool(rows)
        fresh = {}
        for pid, th, title, hp, name, body in rows:
            same = con.execute("select v.blob, v.dim from vec v join paras q on q.id=v.para_rowid where q.thash=? and v.model=? limit 1", (th, MODEL)).fetchone()
            if same:
                con.execute("insert or replace into vec values (?,?,?,?)", (pid, MODEL, same[1], same[0]))
                reused += 1
            else:
                fresh.setdefault(th, (ctx_text(title, hp, name) + "\n" + body, []))[1].append(pid)
        keys, before = list(fresh), done + reused
        for i in range(0, len(keys), batch):
            part = keys[i:i + batch]
            vecs = vec_of(list(emb.embed([fresh[k][0] for k in part], batch_size=batch)))
            for k, v in zip(part, vecs):
                done += 1
                reused += len(fresh[k][1]) - 1   # the same text in several places is embedded once
                for pid in fresh[k][1]:
                    con.execute("insert or replace into vec values (?,?,?,?)", (pid, MODEL, len(v), v.tobytes()))
        con.commit()
        if done + reused == before:   # nothing stored: stop rather than loop
            return done, reused, True


def model_failure(con):
    """The remembered failed model download (dict with at, host, error) while it is younger than MODEL_RETRY_HOURS, else None."""
    try:
        fail = json.loads(get_meta(con, "model_fail", "null"))
        age = datetime.datetime.now(datetime.timezone.utc) - datetime.datetime.strptime(fail["at"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=datetime.timezone.utc)
    except (ValueError, TypeError, KeyError):
        return None
    return fail if age < datetime.timedelta(hours=MODEL_RETRY_HOURS) and fail.get("host") == hub_host() else None   # another endpoint is a new try


def model_failure_reason(fail):
    return (f"model download from {fail['host']} failed ({fail['error']}) at {fail['at']}; keyword search only. "
            f"A refresh tries again after {MODEL_RETRY_HOURS} h; `embed` tries now")


def embed_cmd(con, deadline=None, batch=32, download=True, force=False):
    """Compute vectors. A download that failed is remembered for MODEL_RETRY_HOURS: inside that window `refresh` makes no
    network attempt (force=True, which the `embed` command passes, tries anyway)."""
    fail = None if force or not download or find_model_file() else model_failure(con)
    if fail:
        return {"embedded": 0, "reused": 0, "pending": None, "model": None, "mode_reason": model_failure_reason(fail)}
    try:
        emb = get_embedder(download=download)
    except ModelUnavailable as e:
        reason = str(e)
        if isinstance(e, ModelDownloadFailed):
            fail = {"at": now_iso(), "host": e.host, "error": e.error}
            set_meta(con, "model_fail", json.dumps(fail))
            reason = model_failure_reason(fail)
        add_note(con, str(e))
        con.commit()
        return {"embedded": 0, "reused": 0, "pending": None, "model": None, "mode_reason": reason}
    done, reused, more = embed_pending(con, emb, batch, deadline)
    set_meta(con, "model", MODEL)
    con.execute("delete from meta where k='model_fail'")
    con.commit()
    return {"embedded": done, "reused": reused, "pending": more, "model": MODEL, "mode_reason": None}


# ───────────────────────── sideload: model import and export ─────────────────────────
MODEL_CACHE_DIR = "models--Qdrant--bge-small-en-v1.5-onnx-Q"   # the hub cache folder fastembed writes for MODEL
MODEL_SIDECARS = {                                              # the other files of the pinned revision and their SHA-256 (read from the cached pinned model)
    "config.json": "13582bcf2effc85b7bf3d3f5532e686bc1c9ce86bb009d10f0ec33cbe92299dd",
    "special_tokens_map.json": "5d5b662e421ea9fac075174bb0688ee0d9431699900b90662acd44b2a350503a",
    "tokenizer.json": "d241a60d5e8f04cc1b2b3e9ef7a4921b27bf526d9f6050ab90f9267a1f9e5c66",
    "tokenizer_config.json": "0b29c7bfc889e53b36d9dd3e686dd4300f6525110eaa98c76a5dafceb2029f53",
}
MODEL_MAX_BYTES = 400 << 20                                     # the model is about 66 MB: a bigger package is not it
MODEL_IMPORT_HINT = "doc_index.py model import <folder or .zip with the model>"


class SideloadError(Exception):
    """A model import or export that cannot go ahead: the message is the one line shown."""


def _sha_of(f):
    h = hashlib.sha256()
    for chunk in iter(lambda: f.read(1 << 20), b""):
        h.update(chunk)
    return h.hexdigest()


def _below_model_dir(parts):
    """The path inside the model's cache folder (snapshots/<revision>/<file> or refs/<name>) for a path in a package, else
    None. The cache folder may sit at any depth, or be missing when the package is that folder's contents."""
    for i, p in enumerate(parts):
        if fnmatch.fnmatch(p, MODEL_GLOB):
            rest = parts[i + 1:]
            return tuple(rest) if len(rest) >= 2 and rest[0] in ("snapshots", "refs") else None
    for i, p in enumerate(parts):
        if p in ("snapshots", "refs"):
            return tuple(parts[i:]) if len(parts) - i >= 2 else None
    return None


def _find_model_folder(src, depth=3):
    todo = [(src, 0)]
    while todo:
        d, k = todo.pop(0)
        if fnmatch.fnmatch(os.path.basename(d), MODEL_GLOB) or (k == 0 and os.path.isdir(os.path.join(d, "snapshots"))):
            return d
        if k < depth:
            try:
                todo += [(e.path, k + 1) for e in sorted(os.scandir(d), key=lambda e: e.name) if e.is_dir(follow_symlinks=False)]
            except OSError:
                pass
    return None


@contextlib.contextmanager
def model_sources(src):
    """{path inside the cache folder: (size, opener)} for the files of the model package at `src`, a folder or a .zip. The
    hub's `blobs` copy is left out (the snapshot files are the real ones). SideloadError when `src` is neither, or a zip
    entry would leave its folder."""
    src, zf, out = os.path.abspath(src), None, {}
    try:
        if os.path.isdir(src):
            top = _find_model_folder(src)
            if top is None:
                raise SideloadError(f"no model folder ({MODEL_GLOB}) found in {src}")
            for sub in ("snapshots", "refs"):
                for d, _dirs, files in os.walk(os.path.join(top, sub)):
                    for n in files:
                        full = os.path.join(d, n)
                        out[tuple(os.path.relpath(full, top).replace("\\", "/").split("/"))] = (os.path.getsize(full), lambda p=full: open(p, "rb"))
        elif os.path.isfile(src) and zipfile.is_zipfile(src):
            zf = zipfile.ZipFile(src)
            for info in zf.infolist():
                name = info.filename.replace("\\", "/")
                parts = name.rstrip("/").split("/")
                if name.startswith("/") or any(p in ("", ".", "..") or ":" in p for p in parts):
                    raise SideloadError(f"{os.path.basename(src)} holds an unsafe path ({info.filename!r}); nothing was imported")
                if (info.external_attr >> 16) & 0o170000 == 0o120000:      # a symbolic link is stored as its target path, not as a file
                    raise SideloadError(f"{os.path.basename(src)} holds a symbolic link ({info.filename!r}); nothing was imported")
                rel = None if info.is_dir() else _below_model_dir(parts)
                if rel:
                    out[rel] = (info.file_size, lambda i=info: zf.open(i))
        else:
            raise SideloadError(f"{src} is not a folder or a .zip file")
        yield out
    finally:
        if zf:
            zf.close()


def _copy_hashed(f, dest):
    """Copy the open file `f` to `dest`; the SHA-256 of the bytes that were copied."""
    h = hashlib.sha256()
    with open(dest, "wb") as g:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
            g.write(chunk)
    return h.hexdigest()


def _model_import(src):
    with model_sources(src) as files:
        revs = sorted({r[1] for r in files if len(r) == 3 and r[0] == "snapshots" and r[2] == MODEL_FILE})
        if not revs:
            raise SideloadError(f"{MODEL_FILE} is missing from {src} (expected {MODEL_CACHE_DIR}/snapshots/<revision>/{MODEL_FILE})")
        rev = got = None
        for r in revs:
            with files[("snapshots", r, MODEL_FILE)][1]() as f:
                got = _sha_of(f)
            if got == MODEL_SHA256:
                rev = r
                break
        if rev is None:
            raise SideloadError(f"{MODEL_FILE} in {src} does not match the pinned SHA-256 (got {got[:12]}..., want {MODEL_SHA256[:12]}...); nothing was imported")
        pins = {MODEL_FILE: MODEL_SHA256, **MODEL_SIDECARS}
        gone = [n for n in MODEL_SIDECARS if ("snapshots", rev, n) not in files]
        if gone:
            raise SideloadError(f"{' and '.join(gone)} missing beside {MODEL_FILE} in {src}; nothing was imported")
        keep = {("snapshots", rev, n): files[("snapshots", rev, n)] for n in pins}      # only the pinned files are taken, whatever else the package holds
        for n, want in MODEL_SIDECARS.items():
            with keep[("snapshots", rev, n)][1]() as f:
                if _sha_of(f) != want:
                    raise SideloadError(f"{n} in {src} does not match its pinned SHA-256; nothing was imported")
        total = sum(v[0] for v in keep.values())
        if total > MODEL_MAX_BYTES:
            raise SideloadError(f"{src} holds {total >> 20} MB of model files, more than the {MODEL_MAX_BYTES >> 20} MB the model needs; nothing was imported")
        final, stage = os.path.join(models_dir(), MODEL_CACHE_DIR), os.path.join(models_dir(), f".import-{os.getpid()}")
        shutil.rmtree(stage, ignore_errors=True)
        try:
            top = os.path.join(stage, MODEL_CACHE_DIR)
            for r, (_size, opener) in keep.items():
                dest = os.path.join(top, *r)
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                with opener() as f:
                    if _copy_hashed(f, dest) != pins[r[2]]:       # what was copied is what was checked, or the old model stays as it is
                        raise SideloadError(f"{r[2]} in {src} changed while it was being copied; nothing was imported")
            os.makedirs(os.path.join(top, "refs"), exist_ok=True)
            with open(os.path.join(top, "refs", "main"), "w", encoding="utf-8", newline="") as g:
                g.write(rev)                      # the hub's offline lookup resolves the revision through refs/main
            old = os.path.join(stage, "old")
            if os.path.isdir(final):
                os.replace(final, old)
            try:
                os.replace(top, final)
            except OSError:
                if os.path.isdir(old):
                    os.replace(old, final)
                raise
        except OSError as e:
            raise SideloadError(f"cannot write under {models_dir()}: {e}") from e
        finally:
            shutil.rmtree(stage, ignore_errors=True)
    return {"model": MODEL, "dir": final, "files": len(keep) + 1, "bytes": total}


def model_import(src):
    """Put the pinned embedding model from `src` (a folder, or a .zip, laid out as the hub cache that fastembed writes under
    models_dir()) into models_dir(). Every file of the pinned revision (MODEL_FILE and MODEL_SIDECARS) must be there and match
    its SHA-256 before anything is copied, and the copy is checked again as it is made; it is staged beside the final folder
    and moved in whole, so any failure leaves the old model as it was. Other files in the package are ignored, a symbolic
    link or an entry leaving its folder in a zip is refused, and a damaged zip is one clear line. Never touches the network."""
    try:
        return _model_import(src)
    except (zipfile.BadZipFile, zlib.error, RuntimeError, NotImplementedError) as e:    # a bad CRC, bad deflate data, an encrypted entry
        raise SideloadError(f"{src} is damaged or not a plain zip ({type(e).__name__}: {e}); nothing was imported") from e


def model_export(out):
    """Pack the cached embedding model into a .zip that `model import` takes, so a machine that cannot reach the hub can be
    given it. The file is hashed first; a cache that does not match the pin is reported, not deleted, and nothing is written."""
    f = find_model_file()
    if not f:
        raise SideloadError(f"no cached embedding model under {models_dir()}: run `embed` once on a machine that can reach the hub")
    if sha256(f) != MODEL_SHA256:
        raise SideloadError(f"the cached {MODEL_FILE} does not match the pinned SHA-256: nothing exported (run `embed` again to fetch it)")
    snap = os.path.dirname(f)
    top = os.path.dirname(os.path.dirname(snap))
    out = os.path.abspath(out)
    if os.path.isdir(out):
        raise SideloadError(f"{out} is a folder, not a .zip file name")
    tmp, n = out + ".tmp", 0
    try:
        os.makedirs(os.path.dirname(out), exist_ok=True)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
            for base in (os.path.join(top, "refs"), snap):
                for d, dirs, files in os.walk(base):
                    dirs.sort()
                    for name in sorted(files):
                        full = os.path.join(d, name)
                        zf.write(full, MODEL_CACHE_DIR + "/" + os.path.relpath(full, top).replace("\\", "/"))
                        n += 1
        os.replace(tmp, out)
    except OSError as e:
        raise SideloadError(f"cannot write {out}: {e}") from e
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return {"model": MODEL, "path": out, "files": n, "bytes": os.path.getsize(out)}


def keyword_reason(con, nvec, mf):
    """Why search is keyword-only right now (everything in the way, in one sentence), or None when it is hybrid. Reads the
    disk and the index only: no model load, no network."""
    why = []
    gone = [m for m in ("numpy", "fastembed") if importlib.util.find_spec(m) is None]
    if gone:
        why.append(f"{' and '.join(gone)} not installed in this Python")
    if not mf:
        fail = model_failure(con)
        why.append("offline (HF_HUB_OFFLINE is set) and the embedding model is not downloaded" if offline() else
                   f"the embedding model is not downloaded (the download from {fail['host']} failed: {fail['error']})" if fail else
                   "the embedding model is not downloaded")
    elif not nvec:
        why.append("no vectors yet: run `embed` (or `refresh`)")
    return "; ".join(why) or None


# ───────────────────────── refresh ─────────────────────────
def try_lock(path):
    """An open, locked file handle (kept until the process ends), or None when another process holds the lock."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    fh = open(path, "a+b")
    try:
        fh.seek(0)
        if os.name == "nt":
            import msvcrt
            msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
        return fh
    except OSError:
        fh.close()
        return None


def lock_path(db):
    return os.path.join(os.path.dirname(os.path.abspath(db)), "refresh.lock")


def find_memcap():
    here = os.path.dirname(os.path.abspath(__file__))
    cands = [os.path.join(here, *[".."] * n, "tools", "memcap", "memcap.py") for n in range(4, 8)]
    cands.append(os.path.join(os.environ.get("CLAUDE_CONFIG_DIR") or os.path.join(os.path.expanduser("~"), ".claude"),
                              "workbench", "kit", "tools", "memcap", "memcap.py"))
    return next((os.path.normpath(c) for c in cands if os.path.isfile(c)), None)


def start_detached_refresh(db, extra=()):
    """Start `refresh --all --max-minutes 10` as a detached process under memcap (2,048 MB). False when memcap is not
    found or a refresh already holds the lock."""
    memcap = find_memcap()
    held = try_lock(lock_path(db))
    if held is not None:
        held.close()
    if held is None or memcap is None:
        return False
    argv = [sys.executable, memcap, "--mb", "2048", "--", sys.executable, os.path.abspath(__file__), "--db", db,
            "refresh", "--all", "--max-minutes", "10", *extra]
    kw = {"creationflags": 0x00000008 | 0x00000200} if os.name == "nt" else {"start_new_session": True}
    subprocess.Popen(argv, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, close_fds=True, **kw)
    return True


def lazy_refresh(con, db):
    if os.environ.get("WB_FIND_LAZY") == "0" or not con.execute("select 1 from roots limit 1").fetchone():
        return False
    last = get_meta(con, "last_refresh")
    if last:
        age = datetime.datetime.now(datetime.timezone.utc) - datetime.datetime.strptime(last, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=datetime.timezone.utc)
        if age < datetime.timedelta(hours=STALE_HOURS):
            return False
    return start_detached_refresh(db)


def refresh(con, db, label=None, rebuild=False, embed="auto", max_minutes=None, log=print):
    held = try_lock(lock_path(db))
    if held is None:
        return {"skipped": "another refresh is running"}
    try:
        deadline = time.monotonic() + max_minutes * 60 if max_minutes else None
        roots = con.execute("select id, label, path from roots" + (" where label=?" if label else "") + " order by id", (label,) if label else ()).fetchall()
        if label and not roots:
            raise SystemExit(f"no root named '{label}'")
        res, partial = {}, False
        for r in roots:
            try:
                res[r[1]] = index_root(con, r, rebuild=rebuild, force=rebuild, deadline=deadline, prune=True, log=log)
            except SystemExit as e:   # a root that is missing or offline must not stop the others
                res[r[1]] = {"error": str(e)}
                continue
            partial = partial or res[r[1]]["partial"]
        emb_done = True
        out = {"roots": res, "embed": None}
        if embed != "off" and not partial and not (deadline and time.monotonic() > deadline):
            out["embed"] = embed_cmd(con, deadline, download=True)
            emb_done = not out["embed"]["pending"]
            if out["embed"]["mode_reason"]:
                log("embedding: " + out["embed"]["mode_reason"])
        if not partial and emb_done:
            set_meta(con, "last_refresh", now_iso())
            con.commit()
        return out
    finally:
        held.close()


# ───────────────────────── search ─────────────────────────
def fts_terms(q):
    terms = re.findall(r"\w+", q.lower())
    return [t for t in terms if t not in STOP] or terms


def snippet_of(body, terms, width=240):
    flat = " ".join(body.split())
    low = flat.lower()
    pos = min((low.find(t) for t in terms if len(t) > 2 and t in low), default=0)
    a = max(0, pos - width // 3)
    s = ("..." if a else "") + flat[a:a + width] + ("..." if a + width < len(flat) else "")
    ts = [re.escape(t) for t in terms if len(t) > 2]
    return re.sub(r"(?i)\b(" + "|".join(ts) + r")\w*", lambda m: "[" + m.group(0) + "]", s, count=12) if ts else s


def scope_labels(con, label):
    """None without a --root; else [label] plus the labels of every root nested under it: a root owns the roots inside its
    folder (their files are not indexed under it), so searching it covers them."""
    if not label:
        return None
    row = con.execute("select id, path from roots where label=?", (label,)).fetchone()
    if row is None:
        return [label]
    inner = nested_roots(con, row[0], row[1])
    return [label] + [lab for lab, p in con.execute("select label, path from roots where id<>? order by id", (row[0],)) if norm_path(p) in inner]


def filters_sql(labels, path, source):
    sql, args = "", []
    if labels:
        sql += " and r.label in (%s)" % ",".join("?" * len(labels))
        args.extend(labels)
    if path:
        sql += " and instr(lower(f.ref), ?) > 0"
        args.append(path.lower())
    if source:
        sql += " and lower(coalesce(f.source, '')) = ?"
        args.append(source.lower())
    return sql, args


def bm25_arm(con, terms, labels, path, source):
    if not terms:
        return []
    fsql, fargs = filters_sql(labels, path, source)
    for op in (" AND ", " OR "):   # every word first; any word only when nothing matches them all
        q = op.join('"' + t.replace('"', "") + '"' for t in terms)
        rows = con.execute("select fts.rowid from fts join paras p on p.id=fts.rowid join files f on f.id=p.file_id "
                           "join roots r on r.id=f.root_id where fts match ?" + fsql + " order by bm25(fts) limit ?",
                           [q, *fargs, TOPK]).fetchall()
        if rows:
            return [r[0] for r in rows]
    return []


def dense_arm(con, qvec, labels, path, source):
    # ponytail: every vector is read per query (about half a second per 100,000 paragraphs); cache a matrix file or add
    # sqlite-vec when that stops being fast enough.
    import numpy as np
    fsql, fargs = filters_sql(labels, path, source)
    rows = con.execute("select v.para_rowid, v.blob from vec v join paras p on p.id=v.para_rowid join files f on f.id=p.file_id "
                       "join roots r on r.id=f.root_id where v.model=?" + fsql, [MODEL, *fargs]).fetchall()
    if not rows:
        return []
    ids = np.fromiter((r[0] for r in rows), dtype=np.int64, count=len(rows))
    mat = np.frombuffer(b"".join(r[1] for r in rows), dtype=np.float32).reshape(len(rows), -1)
    sims = mat @ qvec
    top = np.argsort(-sims, kind="stable")[:TOPK]
    return [int(ids[i]) for i in top]


def search(con, db, query, mode="auto", limit=10, root=None, path=None, source=None, rerank=False, lazy=True, per_file=3):
    terms = fts_terms(query)
    nvec = con.execute("select count(*) from vec where model=?", (MODEL,)).fetchone()[0]
    reason, emb = None, None
    if mode == "keyword":
        reason = "keyword mode requested"
    elif nvec == 0:
        fail = None if find_model_file() else model_failure(con)
        reason = model_failure_reason(fail) if fail else "no vectors yet: run `embed` (or `refresh`) while online"
    else:
        try:
            emb = get_embedder(download=False)
        except ModelUnavailable as e:
            fail = model_failure(con) if isinstance(e, ModelMissing) else None   # a corrupt model keeps its own reason
            reason = model_failure_reason(fail) if fail else str(e)
            add_note(con, str(e))
    scope = scope_labels(con, root)
    bm = bm25_arm(con, terms, scope, path, source)
    de = []
    if emb is not None:
        qv = vec_of(next(iter(emb.query_embed(query))))
        de = dense_arm(con, qv, scope, path, source)
    rank = {}
    for arm, ids in (("bm25", bm), ("dense", de)):
        for i, pid in enumerate(ids, 1):
            rank.setdefault(pid, {"bm25": None, "dense": None})[arm] = i
    fused = {pid: sum(1.0 / (RRF_K + r) for r in arms.values() if r) for pid, arms in rank.items()}
    order = sorted(fused, key=lambda p: -fused[p])
    hits, count = [], {}
    if order:
        q = ",".join("?" * len(order))
        info = {r[0]: r for r in con.execute(
            f"select p.id, r.label, f.rel, f.ref, p.n, p.heading, f.title, p.body, f.label, f.url, f.source, f.mtime, p.locator, p.hidden, "
            f"f.fm_status from paras p join files f on f.id=p.file_id join roots r on r.id=f.root_id where p.id in ({q})", order)}
        for pid in sorted(order, key=lambda p: (info[p][14] == "archived", -fused[p])):   # archived items rank below active ones
            _, rlabel, rel, ref, n, hp, title, body, lab, url, src, mt, loc, hid, _st = info[pid]
            if count.get(ref, 0) >= per_file:
                continue
            count[ref] = count.get(ref, 0) + 1
            hits.append({"ref": f"{ref}#p{n}", "root": rlabel, "path": rel, "heading": hp, "title": title,
                         "snippet": snippet_of(body, terms), "score": round(fused[pid], 5), "arms": rank[pid], "label": lab, "url": url,
                         "source": src, "mtime": iso(mt), "locator": loc, "flags": ["hidden-content"] if hid else []})
            if len(hits) >= limit:
                break
    if rerank and hits:
        try:
            hits = rerank_hits(query, hits)
        except Exception as e:   # noqa: BLE001
            reason = ((reason + "; ") if reason else "") + f"rerank unavailable ({type(e).__name__})"
    started = lazy_refresh(con, db) if lazy else False
    cov = coverage(con)
    return {"query": query, "mode": "hybrid" if emb is not None else "keyword", "mode_reason": reason, "refresh_started": started,
            "hits": hits, "coverage": {"roots": [r[0] for r in con.execute("select label from roots order by id") if scope is None or r[0] in scope],
                                       "skipped": cov["skipped"]}}


def rerank_hits(query, hits):
    """Optional MiniLM cross-encoder over the fused top hits; never downloads (offline only)."""
    os.environ["HF_HUB_OFFLINE"] = "1"
    from fastembed.rerank.cross_encoder import TextCrossEncoder
    enc = TextCrossEncoder("Xenova/ms-marco-MiniLM-L-6-v2", cache_dir=models_dir(), lazy_load=False, local_files_only=True)
    scores = list(enc.rerank(query, [h["snippet"] for h in hits]))
    return [h for _s, h in sorted(zip(scores, hits), key=lambda t: -t[0])]


# ───────────────────────── catalog, show, related ─────────────────────────
def files_search(con, words, limit=20, root=None):
    terms = [t.lower() for t in words.split()]
    if not terms:
        return []
    cond = " and ".join("instr(f.hay, ?) > 0" for _ in terms)
    score = " + ".join("(instr(lower(f.name), ?) > 0) * 3 + (instr(lower(f.title), ?) > 0) * 2" for _ in terms)
    scope = scope_labels(con, root)
    args = [x for t in terms for x in (t, t)] + terms + (scope or []) + [limit]
    rows = con.execute(f"select r.label, f.rel, f.ref, f.name, f.title, f.mtime, f.size, f.label, f.cloud_only, f.status, ({score}) sc "
                       f"from files f join roots r on r.id=f.root_id where {cond}" + (" and r.label in (%s)" % ",".join("?" * len(scope)) if scope else "") +
                       " order by sc desc, f.mtime desc limit ?", args).fetchall()
    return [{"ref": r[2], "root": r[0], "path": r[1], "name": r[3], "title": r[4], "mtime": iso(r[5]), "size": r[6], "label": r[7],
             "cloud_only": bool(r[8]), "status": r[9]} for r in rows]


def split_ref(ref):
    base, _, frag = ref.strip().partition("#")
    return base, (int(frag[1:]) if re.fullmatch(r"p\d+", frag) else None)


def find_file(con, ref_or_path):
    """The files row (id, ref, root label, rel, abs path, title) for `<root>/<rel>[#pN]` or an absolute path, else None."""
    base, _ = split_ref(ref_or_path)
    q = ("select f.id, f.ref, r.label, f.rel, r.path, f.title, f.dir from files f join roots r on r.id=f.root_id where ")
    row = con.execute(q + "f.ref=?", (base,)).fetchone()
    if row is None and os.path.isabs(base):
        norm = norm_path(base)
        for label, rpath in con.execute("select label, path from roots order by length(path) desc"):   # the most specific root first
            rp = norm_path(rpath)
            if norm.startswith(rp + os.sep):
                row = con.execute(q + "lower(f.ref)=?", ((label + "/" + norm[len(rp) + 1:].replace(os.sep, "/")).lower(),)).fetchone()
                break
    return row


def show(con, ref, context=2):
    base, n = split_ref(ref)
    row = find_file(con, base)
    if row is None:
        raise SystemExit(f"unknown reference: {ref}")
    fid, fref, rlabel, rel, rpath, title, _d = row
    meta = con.execute("select label, url, source, mtime, status, error, hidden from files where id=?", (fid,)).fetchone()
    rows = con.execute("select n, heading, locator, hidden, body from paras where file_id=? order by n", (fid,)).fetchall()
    lines = [f"source: {os.path.join(rpath, *rel.split('/'))}", f"title: {title}"]
    for k, v in zip(("label", "url", "origin", "modified"), meta[:4]):
        if v:
            lines.append(f"{k}: {iso(v) if k == 'modified' else v}")
    if meta[4] != "ok":
        lines.append(f"not indexed: {meta[4]} {meta[5] or ''}".rstrip())
    sel = [r for r in rows if n is None or n - context <= r[0] <= n + context] if n else rows[:50]
    for pn, hp, loc, hid, body in sel:
        tags = " ".join(x for x in ((f"({loc})" if loc else ""), (f"[hidden-content: {hid}]" if hid else "")) if x)
        lines.append(f"{'>>' if pn == n else '  '} {fref}#p{pn} [{hp}] {tags}\n{body}\n")
    if n and not any(r[0] == n for r in rows):
        raise SystemExit(f"no paragraph p{n} in {fref}")
    return "\n".join(lines)


def neighbors(con, ref, rid_dir=None):
    out = {}
    for dst, kind in con.execute("select dst, kind from edges where src=?", (ref,)):
        out.setdefault(dst, set()).add(kind)
    for src, kind in con.execute("select src, kind from edges where dst=?", (ref,)):
        out.setdefault(src, set()).add(kind)
    for (key,) in con.execute("select key from keys where path=?", (ref,)).fetchall():
        peers = [r[0] for r in con.execute("select path from keys where key=? limit ?", (key, FANOUT_CAP + 2))]
        if len(peers) <= FANOUT_CAP:   # a key seen in more than FANOUT_CAP files makes no co-mention edges
            for p in peers:
                out.setdefault(p, set()).add(key.split(":", 1)[0])
    if rid_dir is not None:
        sibs = [r[0] for r in con.execute("select ref from files where root_id=? and dir=? limit ?", (*rid_dir, FANOUT_CAP + 2))]
        if len(sibs) <= FANOUT_CAP:
            for p in sibs:
                out.setdefault(p, set()).add("folder")
    out.pop(ref, None)
    return out


def related(con, ref_or_path, depth=1, limit=20):
    row = find_file(con, ref_or_path)
    if row is None:
        raise SystemExit(f"unknown file: {ref_or_path}")
    start = row[1]
    seen, frontier, found = {start}, [start], []
    for d in range(1, max(1, depth) + 1):
        nxt = []
        for ref in frontier:
            r = con.execute("select root_id, dir from files where ref=?", (ref,)).fetchone()
            for p, kinds in neighbors(con, ref, (r[0], r[1] or "") if r else None).items():
                if p not in seen:
                    seen.add(p)
                    nxt.append(p)
                    found.append((d, p, sorted(kinds)))
        frontier = nxt
    meta = {r[0]: r for r in con.execute("select ref, rel, title, mtime from files where ref in (%s)" % ",".join("?" * len(seen)), list(seen))}
    found = [f for f in found if f[1] in meta]
    found.sort(key=lambda f: (f[0], -len(f[2]), -(meta[f[1]][3] or 0), f[1]))
    return {"start": start, "related": [{"ref": p, "path": meta[p][1], "title": meta[p][2], "via": via, "depth": d, "mtime": iso(meta[p][3])}
                                        for d, p, via in found[:limit]]}


def newest(con, what):
    row = find_file(con, what) if "/" in what.split("#")[0] or os.path.isabs(what) else None
    if row is None:
        hit = files_search(con, what, 1)
        row = find_file(con, hit[0]["ref"]) if hit else None
    if row is None:
        raise SystemExit(f"nothing found for: {what}")
    ref, chain = row[1], []
    while True:
        up = con.execute("select src from edges where dst=? and kind='supersedes' limit 1", (ref,)).fetchone()
        if not up:
            break
        chain.append(ref)
        ref = up[0]
    f = con.execute("select ref, rel, title, mtime from files where ref=?", (ref,)).fetchone()
    return {"query": what, "newest": {"ref": f[0], "path": f[1], "title": f[2], "mtime": iso(f[3])}, "superseded": chain}


# ───────────────────────── discover, status, purge ─────────────────────────
def count_files(folder, cap=5000):
    on, off, n = 0, 0, 0
    for _rel, st in walk(folder):
        n += 1
        on, off = (on, off + 1) if is_cloud_only(st) else (on + 1, off)
        if n >= cap:
            return on, off, True
    return on, off, False


def discover(con=None):
    """Libraries on this machine. Read-only: directory entries only, nothing is opened or hydrated."""
    home = os.path.expanduser("~")
    found = {}

    def add(path, kind, label):
        if os.path.isdir(path):
            found.setdefault(os.path.normcase(os.path.abspath(path)), (os.path.abspath(path), kind, label))
    for env in ("OneDriveCommercial", "OneDrive"):
        v = os.environ.get(env)
        if v and os.path.basename(v).lower().startswith("onedrive -"):
            add(v, "onedrive-business", os.path.basename(v))
    try:
        for name in os.listdir(home):
            if name.lower().startswith("onedrive -"):
                add(os.path.join(home, name), "onedrive-business", name)
    except OSError:
        pass
    if os.name == "nt" and os.environ.get("WB_DISCOVER_REGISTRY") != "0":
        try:
            import winreg
            base = r"Software\SyncEngines\Providers\OneDrive"
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base) as k:
                for i in range(winreg.QueryInfoKey(k)[0]):
                    with winreg.OpenKey(k, winreg.EnumKey(k, i)) as s:
                        try:
                            mount = winreg.QueryValueEx(s, "MountPoint")[0]
                            ltype = str(winreg.QueryValueEx(s, "LibraryType")[0]).lower()
                        except OSError:
                            continue
                        if ltype in ("teamsite", "documentlibrary"):
                            add(mount, "sharepoint-sync", os.path.basename(mount))
        except OSError:
            pass
    add(os.path.join(home, "Documents"), "documents", "Documents")
    libs = []
    for path, kind, label in found.values():
        on, off, sampled = count_files(path)
        libs.append({"path": path, "kind": kind, "label": label, "on_device_files": on, "cloud_only_files": off, "sampled": sampled})
    if con is not None:
        for label, path, kind in con.execute("select label, path, kind from roots where kind in ('captured','howto','kb')"):
            on, off, sampled = count_files(path) if os.path.isdir(path) else (0, 0, False)
            libs.append({"path": path, "kind": kind, "label": label, "on_device_files": on, "cloud_only_files": off, "sampled": sampled})
    return {"libraries": libs}


def status(con, db):
    roots = []
    for rid, label, path in con.execute("select id, label, path from roots order by id"):
        vec = con.execute("select count(*) from vec v join paras p on p.id=v.para_rowid join files f on f.id=p.file_id "
                          "where f.root_id=?", (rid,)).fetchone()[0]
        roots.append({"label": label, "path": path, "vectors": vec,
                      "files": dict(con.execute("select status, count(*) from files where root_id=? group by status", (rid,)))})
    nvec = con.execute("select count(*) from vec").fetchone()[0]
    paras = con.execute("select count(*) from paras").fetchone()[0]
    mf = find_model_file()
    problems = [{"ref": r[0], "status": r[1], "reason": r[2], "error": r[3]} for r in con.execute(
        "select ref, status, reason, error from files where status in ('failed','skipped','labelled','cloud-only') order by status, ref limit 200")]
    why = keyword_reason(con, nvec, mf)
    return {"db": db, "schema_version": con.execute("pragma user_version").fetchone()[0],
            "migrated_from": json.loads(get_meta(con, "migrated_from", "[]")), "roots": roots, "coverage": coverage(con),
            "paragraphs": paras, "vectors": nvec, "model": MODEL if (nvec or mf) else None,
            "mode": "keyword" if why else "hybrid", "keyword_reason": why, "model_import": None if mf else MODEL_IMPORT_HINT,
            "last_refresh": get_meta(con, "last_refresh"),
            "notes": json.loads(get_meta(con, "notes", "[]")), "problems": problems}


def purge(db, log=print):
    """Delete the index file and anything beside it that holds copies (the 0.2.0 plaintext cache folder)."""
    import shutil
    removed = []
    for f in (db, db + "-wal", db + "-shm", os.path.join(os.path.dirname(os.path.abspath(db)), "refresh.lock")):
        if os.path.exists(f):
            os.remove(f)
            removed.append(f)
    cache = os.path.join(os.path.dirname(os.path.abspath(db)), "cache")
    if os.path.isdir(cache):
        shutil.rmtree(cache)
        removed.append(cache)
    try:
        os.rmdir(os.path.dirname(os.path.abspath(db)))   # only if now empty
    except OSError:
        pass
    log(f"purged {len(removed)} item(s) for index {db}" if removed else f"nothing to purge at {db}")
    return removed


def purge_root(con, label, drop_root=False):
    row = con.execute("select id from roots where label=?", (label,)).fetchone()
    if not row:
        raise SystemExit(f"no root named '{label}'")
    purge_files(con, "root_id=?", (row[0],))
    con.execute("delete from files where root_id=?", (row[0],))
    if drop_root:
        con.execute("delete from roots where id=?", (row[0],))
    con.commit()
    con.execute("pragma wal_checkpoint(truncate)")
    con.execute("vacuum")


# ───────────────────────── command line ─────────────────────────
def reexec_if_needed(cmd):
    """Heavy commands re-run themselves under config.json "knowledge_python" when a needed module is missing here."""
    need = {"index": ["markitdown"], "refresh": ["markitdown"], "convert": ["markitdown"], "embed": ["numpy", "fastembed"],
            "search": ["numpy", "fastembed"], "status": ["numpy", "fastembed"], "report": ["numpy", "fastembed"]}.get(cmd, [])
    py = wb_config().get("knowledge_python")
    if need and py and os.path.isfile(py) and os.environ.get("WB_KNOWLEDGE_REEXEC") != "1" and \
            os.path.normcase(os.path.abspath(py)) != os.path.normcase(os.path.abspath(sys.executable)) and \
            any(importlib.util.find_spec(m) is None for m in need):
        sys.exit(subprocess.call([py, os.path.abspath(__file__), *sys.argv[1:]], env={**os.environ, "WB_KNOWLEDGE_REEXEC": "1"}))


def out_json(obj):
    print(json.dumps(obj, ensure_ascii=False, indent=1))


def table(rows, right=(), indent=""):
    """Aligned text lines: every column is as wide as its longest cell plus two spaces, whatever the data (a long label
    never runs into the next column). Columns listed in `right` are right-aligned; the last column is not padded."""
    w = [max(len(str(c)) for c in col) for col in zip(*rows)]
    return [(indent + "  ".join(str(c).rjust(w[i]) if i in right else str(c).ljust(w[i]) for i, c in enumerate(r))).rstrip() for r in rows]


def main(argv=None):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:   # noqa: BLE001
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", help="index file (default: the shared knowledge index)")
    ap.add_argument("--folder", dest="top_folder", help="use this folder's own index file")
    ap.add_argument("--purge", action="store_true", help="delete the index (or one root's content with --root)")
    ap.add_argument("--root", dest="purge_root")
    sub = ap.add_subparsers(dest="cmd")

    def cmd(name, *args, **flags):
        p = sub.add_parser(name)
        for a in args:
            p.add_argument(a)
        for f, kw in flags.items():
            p.add_argument("--" + f.replace("_", "-"), **kw)
        p.add_argument("--folder", dest="sub_folder", help=argparse.SUPPRESS)
        return p
    J = {"action": "store_true"}
    cmd("discover", json=J)
    cmd("add-root", "path", label={}, json=J)
    cmd("remove-root", "label", purge=J)
    cmd("roots", json=J)
    cmd("index", "folder", rebuild=J, force=J, include_labelled=J, max_mb={"type": int, "default": 50}, max_unzipped_mb={"type": int, "default": 200})
    cmd("refresh", all=J, root={}, rebuild=J, embed={"choices": ["auto", "off"], "default": "auto"}, max_minutes={"type": float}, detach=J, json=J)
    cmd("embed", model={"default": MODEL}, batch={"type": int, "default": 32}, max_minutes={"type": float}, json=J)
    cmd("search", "query", mode={"choices": ["auto", "hybrid", "keyword"], "default": "auto"}, limit={"type": int, "default": 10},
        root={}, path={}, source={}, rerank=J, json=J)
    cmd("show", "ref", context={"type": int, "default": 2})
    cmd("files", "words", limit={"type": int, "default": 20}, root={}, json=J)
    cmd("related", "target", depth={"type": int, "default": 1}, limit={"type": int, "default": 20}, json=J)
    cmd("newest", "target", json=J)
    cmd("convert", "file", include_labelled=J, json=J)
    cmd("label", "file", json=J)
    cmd("model", "action", "path", json=J)
    cmd("status", json=J)
    cmd("report", json=J)
    a = ap.parse_args(argv)
    if not a.cmd and not a.purge:
        ap.error("a command is required")
    reexec_if_needed(a.cmd)
    db = resolve_db(a.db, getattr(a, "sub_folder", None) or a.top_folder)
    try:
        return run(a, db)
    except SystemExit:
        raise
    except Exception as e:   # noqa: BLE001
        print(f"error: {type(e).__name__}: {e}", file=sys.stderr)
        return 1


def run(a, db):
    if a.purge and not a.cmd:
        if a.purge_root:
            con = connect(db)
            purge_root(con, a.purge_root)
            con.close()
            print(f"purged content of root {a.purge_root}")
        else:
            purge(db)
        return 0
    if a.cmd == "convert":
        d = convert_doc(a.file, a.include_labelled)
        md = "\n\n".join((f"[hidden: {k}]\n" if k else "") + t for t, _l, k in d["segments"])
        out_json({"path": a.file, "label": d["label"], "markdown": md, "hidden_content": d["hidden_content"], "error": d["error"]})
        return 0
    if a.cmd == "label":
        out_json({"path": a.file, "label": file_label(a.file)})
        return 0
    if a.cmd == "model":
        if a.action not in ("import", "export"):
            print("usage: doc_index.py model import <folder or .zip> | model export <out.zip>", file=sys.stderr)
            return 2
        try:
            res = (model_import if a.action == "import" else model_export)(a.path)
        except SideloadError as e:
            print(f"error: model {a.action}: {e}", file=sys.stderr)
            return 1
        if a.json:
            out_json(res)
        elif a.action == "import":
            print(f"imported the embedding model ({res['files']} files, {res['bytes'] >> 20} MB) into {res['dir']}; next: doc_index.py embed")
        else:
            print(f"exported the embedding model ({res['files']} files, {res['bytes'] >> 20} MB) to {res['path']}; on the other machine: doc_index.py model import \"{res['path']}\"")
        return 0
    con = connect(db)
    try:
        if a.cmd == "discover":
            d = discover(con)
            if a.json:
                out_json(d)
            else:
                for L in d["libraries"]:
                    print(f"{L['kind']:<18}{L['label']}  {L['path']}  on-device {L['on_device_files']}  cloud-only {L['cloud_only_files']}{'+' if L['sampled'] else ''}")
        elif a.cmd == "add-root":
            label = add_root(con, a.path, a.label)
            r = {"root": label, "path": os.path.abspath(a.path)}
            out_json(r) if a.json else print(f"root {label}: {r['path']}")
        elif a.cmd == "remove-root":
            purge_root(con, a.label, drop_root=True)
            print(f"removed root {a.label}")
        elif a.cmd == "roots":
            rows = [{"label": l, "path": p, "kind": k, "files": n} for l, p, k, n in con.execute(
                "select r.label, r.path, r.kind, (select count(*) from files f where f.root_id=r.id) from roots r order by r.id")]
            out_json({"roots": rows}) if a.json else [print(ln) for ln in table([(r["label"], r["kind"], r["files"], r["path"]) for r in rows], right=(2,))]
        elif a.cmd == "index":
            label = add_root(con, a.folder)
            root = con.execute("select id, label, path from roots where label=?", (label,)).fetchone()
            stats = index_root(con, root, a.rebuild, a.force, a.include_labelled, a.max_mb, a.max_unzipped_mb)
            print(index_stats_line(stats, db))
        elif a.cmd == "refresh":
            if a.detach:
                ok = start_detached_refresh(db, (["--root", a.root] if a.root else []) + (["--rebuild"] if a.rebuild else []))
                out_json({"started": ok}) if a.json else print("refresh started in the background" if ok else "refresh not started (already running, or memcap not found)")
                return 0
            res = refresh(con, db, a.root, a.rebuild, a.embed, a.max_minutes, log=lambda m: print(m, file=sys.stderr))   # notes to stderr, the result to stdout
            out_json(res) if a.json else print(json.dumps(res, ensure_ascii=False))
        elif a.cmd == "embed":
            deadline = time.monotonic() + a.max_minutes * 60 if a.max_minutes else None
            res = embed_cmd(con, deadline, a.batch, force=True)   # asked for by hand: try now, whatever was remembered
            out_json(res) if a.json else print(json.dumps(res))
        elif a.cmd == "search":
            res = search(con, db, a.query, a.mode, a.limit, a.root, a.path, a.source, a.rerank)
            if a.json:
                out_json(res)
            else:
                for h in res["hits"]:
                    extra = " ".join(x for x in ((f"({h['locator']})" if h["locator"] else ""), (f"label={h['label']}" if h["label"] else ""),
                                                 (f"url={h['url']}" if h["url"] else ""), ("[hidden-content]" if h["flags"] else "")) if x)
                    print(f"{h['ref']}  [{h['heading'] or '-'}] {extra}\n    {h['snippet']}")
                if not res["hits"]:
                    print("no matches (try other words; keyword search misses paraphrases until vectors exist)")
                print(f"mode: {res['mode']}" + (f" ({res['mode_reason']})" if res["mode_reason"] else "") +
                      (f"\nnot searched: {json.dumps(res['coverage']['skipped'])}" if res["coverage"]["skipped"] else ""))
        elif a.cmd == "show":
            print(show(con, a.ref, a.context))
        elif a.cmd == "files":
            rows = files_search(con, a.words, a.limit, a.root)
            if a.json:
                out_json({"query": a.words, "files": rows})
            else:
                for r in rows:
                    print(f"{r['ref']}  {r['mtime'] or '-'}  {r['size']}  {r['title']}{'  [cloud-only]' if r['cloud_only'] else ''}")
        elif a.cmd == "related":
            res = related(con, a.target, a.depth, a.limit)
            out_json(res) if a.json else [print(f"d{r['depth']} {','.join(r['via']):<20}{r['ref']}") for r in res["related"]]
        elif a.cmd == "newest":
            res = newest(con, a.target)
            out_json(res) if a.json else print(res["newest"]["ref"])
        else:   # status, report
            st = status(con, db)
            if a.json:
                out_json(st)
            else:
                print(f"index: {db}\nroots {len(st['roots'])}  files {st['coverage']['catalogued']}  indexed {st['coverage']['content_indexed']}  "
                      f"paragraphs {st['paragraphs']}  vectors {st['vectors']}  mode {st['mode']}  last refresh {st['last_refresh'] or 'never'}")
                if st["keyword_reason"]:
                    print(f"  search is keyword-only: {st['keyword_reason']}")
                if st["model_import"]:
                    print(f"  without the internet, add the model from a folder or .zip: {st['model_import']}"
                          "  (build the .zip where the hub is reachable: doc_index.py model export <out.zip>)")
                for ln in table([(r["label"], f"{sum(r['files'].values())} files", f"{r['vectors']} vectors", r["path"]) for r in st["roots"]], right=(1, 2), indent="  "):
                    print(ln)
                for k, n in st["coverage"]["skipped"].items():
                    print(f"  not searched: {k} {n}")
                for n in st["notes"]:
                    print(f"  note: {n}")
                for ln in table([(p["status"], p["ref"], p["error"] or "") for p in st["problems"]], indent="  "):
                    print(ln)
                for old in st["migrated_from"]:
                    print(f"  0.2.0 index kept at {old}; delete it with: doc_index.py --db \"{old}\" --purge")
    finally:
        con.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())

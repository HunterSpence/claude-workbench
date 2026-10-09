#!/usr/bin/env python3
"""make_corpus.py: builds the fictional golden corpus (docx, pptx, xlsx, pdf, md) used by eval_find.py.

  python make_corpus.py DEST        writes the files into DEST (created if needed) and prints the file count

Every document is invented and neutral. Names are slug + extension, and golden.json refers to them by that name.
The corpus is deliberately small and ships with near-topic neighbours (parking vs visitor badges, incident response
vs post-incident review, and so on), so a wrong answer can sit close to the right one.
"""
import os
import sys

# slug, format, title, paragraphs
DOCS = [
    ("parking-permit", "docx", "Parking permit renewal", [
        "Employees who drive to the campus must renew their parking permit once every twelve months. Renewal is handled at the front desk during opening hours.",
        "Bring the expired sticker with you. A fresh sticker for the windscreen is issued on the spot and the old one is scraped off by the facilities team.",
        "If a permit is lost or stolen, report it to the front desk and sign the replacement form. The old permit number is cancelled immediately."]),
    ("visitor-badge", "md", "Visitor badges", [
        "Guests must sign in at reception and wear a lanyard with a printed name tag while they are inside the building.",
        "The host collects the guest from the lobby, stays with them throughout the visit, and returns the tag at the desk on departure."]),
    ("bicycle-storage", "md", "Bicycle storage and showers", [
        "Cyclists can leave bikes in the covered rack behind building two. The cage is locked at night and opened again at six.",
        "Showers and lockers are on the ground floor of the sports annex. Towels are not provided."]),
    ("expense-claims", "docx", "Expense claims", [
        "Submit receipts for reimbursement within thirty days of purchase. Claims without a receipt are returned to the sender.",
        "Managers approve claims below the team limit; the finance office approves the rest. Payment follows in the next monthly run."]),
    ("travel-booking", "pptx", "Booking business travel", [
        "Book flights and hotels through the approved travel portal. Itineraries must be approved before ticketing.",
        "Economy class is the default for trips under six hours. Longer trips may be upgraded with director approval."]),
    ("mileage-allowance", "xlsx", "Mileage allowance", [
        "A private car used for work trips is reimbursed per kilometre at the published rate.",
        "Log the start and end odometer reading for every journey and attach the log to the claim."]),
    ("incident-response", "docx", "Incident response", [
        "When a production service degrades, page the on-call engineer straight away. Severity one means customers cannot complete their core tasks.",
        "The incident commander opens a bridge call and posts a status update every thirty minutes until the service recovers.",
        "Do not make unreviewed changes during an incident. Record every action in the incident channel."]),
    ("post-incident-review", "md", "Post-incident review", [
        "After an outage, hold a blameless review within five working days. The goal is learning, not fault finding.",
        "Write a timeline of events, list the contributing factors, and agree follow-up actions with named owners."]),
    ("maintenance-window", "pdf", "Planned maintenance", [
        "Planned maintenance happens on the first Sunday of each month between two and four in the morning.",
        "Notices go to subscribers one week ahead. Emergency changes outside the window need director approval."]),
    ("password-policy", "docx", "Password policy", [
        "Passphrases must contain at least fourteen characters. Multi-factor sign-in is mandatory for every account.",
        "Reusing a secret across systems is forbidden. A password manager is provided to every employee."]),
    ("disk-encryption", "md", "Laptop disk encryption", [
        "Every laptop must have full disk encryption enabled before it leaves the office.",
        "Recovery keys are escrowed with the security team. Report a stolen machine within one hour so that access can be revoked."]),
    ("phishing-awareness", "pptx", "Spotting phishing", [
        "Do not click links in unexpected messages. Check the sender address before you reply.",
        "Use the report button in the mail client to forward a suspicious message to the security desk. Training simulations run each quarter."]),
    ("annual-leave", "docx", "Annual leave", [
        "Full-time staff accrue two days of paid holiday per month. Up to five unused days carry over into the next calendar year.",
        "Requests are made in the people system at least two weeks ahead and are approved by the line manager."]),
    ("sick-leave", "md", "Sick leave", [
        "When you are unwell, tell your manager before the start of your shift.",
        "A medical certificate is required after three consecutive days of absence."]),
    ("parental-leave", "docx", "Parental leave", [
        "New parents may take up to sixteen weeks of leave, which can be split into blocks within the first year.",
        "Plan the return with your manager and agree a phased schedule for the first month."]),
    ("onboarding-checklist", "xlsx", "Onboarding checklist", [
        "Day one: collect the laptop, building pass and accounts from the service desk.",
        "Week one: meet the team lead, read the team handbook and complete the mandatory safety course."]),
    ("offboarding-checklist", "md", "Offboarding checklist", [
        "When someone leaves, they return all equipment and the building pass on the last day.",
        "Access is revoked at the end of the final day, and open work is handed over to a named colleague."]),
    ("mentoring-program", "pptx", "Mentoring program", [
        "Pair each junior colleague with a senior peer for six months.",
        "Meet every two weeks to discuss goals, feedback and career growth. Pairs report back to the program lead at the end."]),
    ("data-retention", "docx", "Data retention schedule", [
        "Financial records are kept for seven years and then destroyed. Customer correspondence is kept for three years.",
        "A legal hold suspends deletion for the records it names until the hold is lifted in writing."]),
    ("backup-schedule", "xlsx", "Backup schedule", [
        "Nightly snapshots of the file servers are kept for fourteen days. Weekly copies are kept for three months.",
        "Restore tests are performed every quarter and the results are filed with the infrastructure team."]),
    ("archive-requests", "md", "Archive requests", [
        "To retrieve a boxed file from off-site storage, file a request at least two days ahead.",
        "Boxes return to storage after thirty days unless an extension is granted."]),
    ("room-booking", "pdf", "Meeting room booking", [
        "Meeting rooms are reserved in the shared calendar. A booking is released if nobody checks in within ten minutes.",
        "The large boardroom needs approval from the office manager."]),
    ("catering-orders", "docx", "Catering orders", [
        "Order lunch and refreshments through the catering form at least three working days before an event.",
        "State dietary needs such as vegetarian or nut-free so that the kitchen can plan the menu."]),
    ("office-supplies", "md", "Office supplies", [
        "Stationery and printer toner are stocked in the supply cupboard on each floor.",
        "Replenishment orders go in on the first Monday of the month."]),
    ("release-process", "pptx", "Release process", [
        "Freeze the main branch two days before a release. Tag the build and publish the changelog.",
        "Run the smoke checks before announcing the release to customers."]),
    ("code-review", "md", "Code review guidelines", [
        "Every change needs a second pair of eyes before merging. Reviewers check correctness, readability and tests.",
        "Keep changes small enough to read in half an hour. Large changes are split before review."]),
    ("hotfix-procedure", "docx", "Hotfix procedure", [
        "A hotfix skips the normal freeze but still needs one reviewer and a rollback plan.",
        "Backport the fix to the stable branch afterwards and note it in the changelog."]),
    ("vendor-evaluation", "xlsx", "Vendor evaluation", [
        "Score each supplier from one to five on price, quality, support and security.",
        "Shortlist the top three for demonstrations and ask each for two references."]),
    ("contract-renewal", "docx", "Contract renewal", [
        "Review contracts ninety days before expiry. Give notice of non-renewal in writing.",
        "Clauses that renew automatically need a legal review before they are signed."]),
    ("procurement-approval", "md", "Procurement approval", [
        "Purchases above the delegated threshold need two approvals and a purchase order number before the order is placed.",
        "Orders placed without a purchase order number are not paid until the gap is explained."]),
    ("fire-safety", "pdf", "Fire safety", [
        "On hearing the alarm, leave by the nearest marked exit and gather at the assembly point in the car park. Do not use lifts.",
        "Fire wardens check each floor and report to the senior warden. Nobody re-enters until the all-clear."]),
    ("first-aid", "md", "First aid", [
        "First aid kits are mounted beside every stairwell. Trained first aiders wear green lanyards.",
        "Record every use of a kit in the accident book so that supplies can be restocked."]),
    ("ergonomic-workstation", "docx", "Ergonomic workstation", [
        "Adjust the chair so your feet rest flat and the top of the screen sits at eye level.",
        "Take a short break from sitting every hour. Request a sit-stand desk through facilities."]),
    ("volunteer-day", "pptx", "Volunteer day", [
        "Each employee may spend one paid day per year helping a local charity of their choice.",
        "Register the activity with the community team beforehand."]),
    ("charity-matching", "md", "Charity matching", [
        "The company matches personal donations to registered charities up to an annual limit per person.",
        "Send the donation receipt to the community team to trigger the match."]),
    ("recycling-program", "docx", "Recycling program", [
        "Separate paper, glass and plastics into the coloured bins on each floor.",
        "Batteries go to the collection box at reception. Electronics are collected on the last Friday of each month."]),
]


def pdf_bytes(title, paras):
    """A minimal one-page PDF (Helvetica text, one line per wrapped row) that pdfminer can read."""
    lines = [title, ""]
    for p in paras:
        words, row = p.split(), ""
        for w in words:
            if len(row) + len(w) > 78:
                lines.append(row)
                row = ""
            row += (w + " ")
        lines += [row.strip(), ""]
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")   # noqa: E731
    body = "BT /F1 11 Tf 14 TL 50 780 Td " + " ".join(f"({esc(t)}) Tj T*" for t in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(body)} >>\nstream\n{body}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def write(dest, slug, fmt, title, paras):
    path = os.path.join(dest, f"{slug}.{fmt}")
    if fmt == "md":
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"# {title}\n\n" + "\n\n".join(paras) + "\n")
    elif fmt == "docx":
        import docx
        d = docx.Document()
        d.add_heading(title, 1)
        for p in paras:
            d.add_paragraph(p)
        d.save(path)
    elif fmt == "pptx":
        import pptx
        prs = pptx.Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
        for p in paras:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = title
            s.placeholders[1].text = p
        prs.save(path)
    elif fmt == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:30]
        ws.append(["Item", "Detail"])
        for i, p in enumerate(paras, 1):
            ws.append([f"Row {i}", p])
        wb.save(path)
    elif fmt == "pdf":
        with open(path, "wb") as f:
            f.write(pdf_bytes(title, paras))
    else:
        raise ValueError(fmt)
    return os.path.basename(path)


def build(dest):
    os.makedirs(dest, exist_ok=True)
    return [write(dest, *d) for d in DOCS]


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    print(len(build(sys.argv[1])), "files written to", sys.argv[1])
